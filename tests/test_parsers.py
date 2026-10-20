import pytest
from docx import Document

from app.parsers import UnsupportedFormatError, parse_file, parse_text


def test_纯文本解析保留标题层级():
    doc = parse_text("# 登录模块\n用户输入账号密码\n## 异常场景\n密码错误提示")
    assert doc.doc_type == "txt"
    titles = [(s.level, s.title) for s in doc.sections if s.title]
    assert (1, "登录模块") in titles
    assert (2, "异常场景") in titles
    assert "密码错误提示" in doc.full_text


def test_txt文件解析(tmp_path):
    f = tmp_path / "需求.txt"
    f.write_text("充值功能：支持微信与支付宝", encoding="utf-8")
    doc = parse_file(f)
    assert doc.source == "需求.txt"
    assert "充值功能" in doc.full_text


def test_docx解析标题正文与表格(tmp_path):
    path = tmp_path / "需求.docx"
    d = Document()
    d.add_heading("活动系统需求", level=1)
    d.add_paragraph("活动期间每日登录可领取积分。")
    d.add_heading("积分规则", level=2)
    d.add_paragraph("积分上限为 1000 分。")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "等级"
    table.cell(0, 1).text = "积分倍率"
    table.cell(1, 0).text = "VIP1"
    table.cell(1, 1).text = "1.5"
    d.save(str(path))

    doc = parse_file(path)

    assert doc.doc_type == "docx"
    titles = [(s.level, s.title) for s in doc.sections if s.title]
    assert (1, "活动系统需求") in titles
    assert (2, "积分规则") in titles
    assert "积分上限为 1000 分。" in doc.full_text
    assert doc.tables == [[["等级", "积分倍率"], ["VIP1", "1.5"]]]


def test_不限制格式_未知后缀按文本读取_二进制才报不支持(tmp_path):
    f = tmp_path / "需求.xyz"
    f.write_text("# 登录\n输入账号密码登录", encoding="utf-8")
    doc = parse_file(f)
    assert doc.doc_type == "xyz" and "输入账号密码登录" in doc.full_text and doc.sections[0].title == "登录"
    g = tmp_path / "接口.json"
    g.write_text('{"login": "POST /api/login"}', encoding="utf-8")
    assert "POST /api/login" in parse_file(g).full_text
    b = tmp_path / "程序.exe"
    b.write_bytes(b"MZ\x00\x01\x90\x00" * 10)
    with pytest.raises(UnsupportedFormatError, match="不支持"):
        parse_file(b)


def test_excel与csv按表格解析(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "功能清单"
    ws.append(["模块", "功能", "说明"])
    ws.append(["登录", "账号密码登录", "错误 5 次锁定"])
    ws2 = wb.create_sheet("接口")
    ws2.append(["接口", "方法"])
    ws2.append(["/api/login", "POST"])
    x = tmp_path / "需求.xlsx"
    wb.save(x)
    doc = parse_file(x)
    text = doc.full_text
    assert doc.doc_type == "xlsx"
    # 数据型工作表按「列名：值」展开；两列的小表按文档型逐行输出
    assert "# 功能清单" in text and "模块：登录；功能：账号密码登录；说明：错误 5 次锁定" in text
    assert "# 接口" in text and "/api/login ｜ POST" in text
    c = tmp_path / "需求.csv"
    c.write_text("模块,功能\n登录,找回密码\n", encoding="utf-8")
    doc = parse_file(c)
    assert "登录 ｜ 找回密码" in doc.full_text


def test_pptx按页解析文字与图片占位(tmp_path):
    import zipfile

    slide = """<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
      xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
      xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:cSld><p:spTree>
      <p:sp><p:txBody><a:p><a:r><a:t>登录流程</a:t></a:r></a:p><a:p><a:r><a:t>输入账号</a:t></a:r><a:r><a:t>密码</a:t></a:r></a:p></p:txBody></p:sp>
      <p:pic><p:blipFill><a:blip r:embed="rId2"/></p:blipFill></p:pic></p:spTree></p:cSld></p:sld>"""
    rels = """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/></Relationships>"""
    from PIL import Image
    import io
    buf = io.BytesIO(); Image.new("RGB", (8, 8), "white").save(buf, format="PNG")
    f = tmp_path / "需求.pptx"
    with zipfile.ZipFile(f, "w") as zf:
        zf.writestr("ppt/slides/slide1.xml", slide)
        zf.writestr("ppt/slides/_rels/slide1.xml.rels", rels)
        zf.writestr("ppt/media/image1.png", buf.getvalue())
        zf.writestr("ppt/slides/slide2.xml", slide.replace("登录流程", "找回密码").replace('<p:pic><p:blipFill><a:blip r:embed="rId2"/></p:blipFill></p:pic>', ""))
    doc = parse_file(f)
    assert doc.doc_type == "pptx"
    assert doc.sections[0].title == "第 1 页：登录流程" and "输入账号密码" in doc.full_text
    assert len(doc.embedded_images) == 1 and "[[图片:1]]" in doc.full_text
    assert any(s.title == "第 2 页：找回密码" for s in doc.sections)


# ---- 全格式支持：压缩包 / 邮件 / 网页 / 扩展图片 / 老版 Office 转换 / 非文本类型 ----


def _png_bytes(size=(120, 120)):
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, "white").save(buf, format="PNG")
    return buf.getvalue()


def test_压缩包递归解析_文档合并_图片转内嵌占位(tmp_path):
    import io
    import tarfile
    import zipfile

    z = tmp_path / "需求包.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("docs/需求.md", "# 登录\n账号密码登录")
        zf.writestr("docs/清单.csv", "模块,功能\n登录,找回密码\n")
        zf.writestr("原型.png", _png_bytes())
        zf.writestr("__MACOSX/._需求.md", b"junk")
        zf.writestr("demo.mp4", b"\x00\x00\x00 ftypmp42")
    doc = parse_file(z)
    text = doc.full_text
    assert doc.doc_type == "zip"
    assert "文件：docs/需求.md" in text and "账号密码登录" in text and "登录 ｜ 找回密码" in text
    assert len(doc.embedded_images) == 1 and "[[图片:1]]" in text
    assert "demo.mp4" in text and "视频文件" in text and "junk" not in text
    # tar.gz 同样可读
    t = tmp_path / "需求.tar.gz"
    with tarfile.open(t, "w:gz") as tf:
        data = "测试范围：登录与找回密码".encode()
        info = tarfile.TarInfo("范围.txt"); info.size = len(data)
        tf.addfile(info, io.BytesIO(data))
    assert "登录与找回密码" in parse_file(t).full_text


def test_邮件解析_正文与附件(tmp_path):
    from email.message import EmailMessage

    m = EmailMessage()
    m["Subject"] = "登录需求确认"; m["From"] = "pm@x.com"; m["To"] = "qa@x.com"
    m.set_content("请按附件需求设计用例。")
    m.add_attachment("# 需求\n密码错误 5 次锁定".encode(), maintype="text", subtype="markdown", filename="需求.md")
    m.add_attachment(_png_bytes(), maintype="image", subtype="png", filename="原型.png")
    f = tmp_path / "需求.eml"
    f.write_bytes(bytes(m))
    doc = parse_file(f)
    text = doc.full_text
    assert doc.sections[0].title == "邮件：登录需求确认" and "发件人：pm@x.com" in text
    assert "请按附件需求设计用例" in text and "密码错误 5 次锁定" in text
    assert len(doc.embedded_images) == 1


def test_网页文件与扩展图片格式(tmp_path):
    from app.parsers.image import load_image_bytes

    h = tmp_path / "需求.html"
    h.write_text("<html><head><title>登录</title></head><body><h1>规则</h1><p>锁定 30 分钟</p></body></html>", encoding="utf-8")
    doc = parse_file(h)
    assert doc.doc_type == "html" and "# 登录" in doc.full_text and "锁定 30 分钟" in doc.full_text
    from PIL import Image

    b = tmp_path / "截图.bmp"
    Image.new("RGB", (30, 30), "red").save(b, format="BMP")
    data, mime = load_image_bytes(b)
    assert mime == "image/png" and data[:8] == b"\x89PNG\r\n\x1a\n"
    g = tmp_path / "动图.gif"
    Image.new("RGB", (30, 30), "blue").save(g, format="GIF")
    assert load_image_bytes(g)[1] == "image/gif"


def test_老版Office需LibreOffice_未安装给出明确提示(tmp_path, monkeypatch):
    from app.parsers import convert
    from app.parsers.convert import ConversionUnavailableError

    monkeypatch.setattr(convert, "soffice_path", lambda: None)
    for name in ("需求.doc", "清单.xls", "方案.ppt", "文字.wps", "表格.et", "演示.dps", "说明.rtf"):
        f = tmp_path / name
        f.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 64)
        with pytest.raises(ConversionUnavailableError, match="LibreOffice"):
            parse_file(f)
    v = tmp_path / "视频.mp4"
    v.write_bytes(b"\x00" * 16)
    with pytest.raises(UnsupportedFormatError, match="视频文件"):
        parse_file(v)


def test_excel文档型渲染_去空单元格_内嵌截图占位(tmp_path):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XImage
    from PIL import Image

    img = tmp_path / "shot.png"
    Image.new("RGB", (160, 120), "white").save(img)
    wb = Workbook()
    ws = wb.active
    ws.title = "道具系统"
    ws["I2"] = "背包系统"
    ws["I3"] = "1.道具分类按钮，点击可以切换物品显示的种类"
    ws["I4"] = "参考左图"
    ws["I6"] = "道具详情弹窗"
    ws.add_image(XImage(str(img)), "B3")
    f = tmp_path / "需求.xlsx"
    wb.save(f)
    doc = parse_file(f)
    text = doc.full_text
    assert "# 道具系统" in text and "|  |" not in text and " ｜ " not in text
    assert text.index("1.道具分类按钮") < text.index("参考左图") < text.index("道具详情弹窗")
    assert len(doc.embedded_images) == 1 and "[[图片:1]]" in text
    assert text.index("[[图片:1]]") < text.index("道具详情弹窗")


def test_excel数据型渲染_按列名展开(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["模块", "功能", "规则", "优先级"])
    ws.append(["登录", "密码登录", "错误 5 次锁定 30 分钟", "P0"])
    ws.append(["登录", "验证码登录", "60 秒内只能发一次", "P1"])
    f = tmp_path / "清单.xlsx"
    wb.save(f)
    text = parse_file(f).full_text
    assert "模块：登录；功能：密码登录；规则：错误 5 次锁定 30 分钟；优先级：P0" in text
    assert "功能：验证码登录" in text


def test_压缩包不限文件数(tmp_path):
    import zipfile

    z = tmp_path / "大包.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for i in range(1, 602):
            zf.writestr(f"docs/需求{i:03d}.txt", f"第 {i} 条需求")
    doc = parse_file(z)
    text = doc.full_text
    assert "文件：docs/需求001.txt" in text and "文件：docs/需求601.txt" in text and "第 601 条需求" in text


def test_Axure原型导出包_按站点地图逐页提取文字与说明_跳过脚本(tmp_path):
    import json
    import zipfile

    sitemap = {"configuration": {}, "sitemap": {"rootNodes": [
        {"pageName": "首页", "type": "Wireframe", "url": "首页.html", "children": [
            {"pageName": "登录", "type": "Wireframe", "url": "登录.html"}]},
        {"pageName": "设置", "type": "Folder", "children": [{"pageName": "账号安全", "type": "Wireframe", "url": "账号安全.html"}]},
    ]}}
    page = lambda title, *texts: ("<!DOCTYPE html><html><head><title>%s</title><script src='resources/scripts/axure/axQuery.js'></script></head>"
                                  "<body><div id='base'>%s</div><script>$axure.loadCurrentPage((function(){}))</script></body></html>"
                                  % (title, "".join(f"<div id='u{i}' class='ax_default label'><div id='u{i}_div'></div>"
                                                    f"<div id='u{i}_text' class='text'><p><span>{t}</span></p></div></div>" for i, t in enumerate(texts))))
    z = tmp_path / "原型.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("原型/index.html", "<html><body><iframe src='start.html'></iframe></body></html>")
        zf.writestr("原型/start.html", "<html><body>player</body></html>")
        zf.writestr("原型/data/document.js", "$axure.loadDocument(" + json.dumps(sitemap, ensure_ascii=False) + ");")
        zf.writestr("原型/首页.html", page("首页", "欢迎使用", "登录按钮", "登录按钮"))
        zf.writestr("原型/登录.html", page("登录", "请输入手机号", "密码错误 5 次锁定"))
        zf.writestr("原型/账号安全.html", page("账号安全", "修改密码"))
        zf.writestr("原型/files/登录/data.js", '$axure.loadCurrentPage((function() { var _ = function(){}, a="url", b="登录.html", '
                                             'c="<p><span>点击登录后跳转首页</span></p>", d="<p>请输入手机号</p>"; return _(a,b,c,d); })());')
        zf.writestr("原型/resources/scripts/axure/axure.js", "var $axure = {};")
        zf.writestr("原型/resources/css/axure_rp_page.css", "body{}")
        zf.writestr("原型/images/首页/u0.png", _png_bytes())
    doc = parse_file(z)
    text = doc.full_text
    assert doc.doc_type == "zip" and "共 3 个页面" in text
    assert text.index("页面：首页") < text.index("页面：首页 / 登录") < text.index("页面：设置 / 账号安全")
    assert "欢迎使用" in text and "密码错误 5 次锁定" in text and "修改密码" in text
    assert text.count("登录按钮") == 1  # 同页重复文字去重
    assert "点击登录后跳转首页" in text and "【说明 / 备注】" in text
    assert text.count("请输入手机号") == 1  # data.js 中与页面正文相同的文字不重复（正文行带控件编号，比对时去掉）
    assert "[u0] 请输入手机号" in text and "- 点击登录后跳转首页" in text    # 页面行带编号，说明备注不带
    assert "$axure" not in text and "axQuery" not in text and "player" not in text and "u0.png" not in text


def test_普通压缩包内脚本样式资源只汇总不解析(tmp_path):
    import zipfile

    z = tmp_path / "包.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("需求.md", "# 登录\n账号密码登录")
        zf.writestr("app.js", "function login() { return 1; }")
        zf.writestr("style.css", "body { color: red }")
    text = parse_file(z).full_text
    assert "账号密码登录" in text and "function login" not in text and "color: red" not in text
    assert "2 个脚本 / 样式资源文件不参与解析" in text and "app.js" in text


def test_压缩包内文档图片多于外层图片_占位重新编号不连锁覆盖(tmp_path):
    """早先逐张边改边替：docx 第 1 张改成 [[图片:2]] 后，处理 docx 第 2 张时又把它改成 [[图片:3]]……
    连锁覆盖后前面几张图的占位全部丢失、只剩最后几张各重复多次，理解结果随之丢失。"""
    import zipfile

    from docx import Document

    d = tmp_path / "后台需求.docx"
    doc = Document()
    for i in range(1, 4):
        doc.add_paragraph(f"第 {i} 页说明")
        img = tmp_path / f"p{i}.png"
        img.write_bytes(_png_bytes((200, 200)))
        doc.add_picture(str(img))
    doc.save(d)
    z = tmp_path / "需求包.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("封面.png", _png_bytes((200, 200)))
        zf.write(d, "后台需求.docx")
    parsed = parse_file(z)
    placeholders = [s.content for s in parsed.sections if s.content.startswith("[[图片:")]
    assert placeholders == ["[[图片:1]]", "[[图片:2]]", "[[图片:3]]", "[[图片:4]]"]
    assert [img.placeholder for img in parsed.embedded_images] == placeholders


def test_docx表格单元格内图片按位置提取_合并单元格不重复(tmp_path):
    from docx import Document

    d = tmp_path / "表格图.docx"
    doc = Document()
    doc.add_paragraph("页面说明")
    img = tmp_path / "shot.png"; img.write_bytes(_png_bytes((200, 200)))
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).paragraphs[0].add_run().add_picture(str(img))
    t.cell(0, 1).text = "右侧说明"
    t2 = doc.add_table(rows=1, cols=2)
    t2.cell(0, 0).merge(t2.cell(0, 1)).paragraphs[0].add_run().add_picture(str(img))   # 合并单元格：同一图只算一次
    doc.add_paragraph("结尾")
    doc.save(d)
    parsed = parse_file(d)
    contents = [s.content for s in parsed.sections if s.content]
    assert len(parsed.embedded_images) == 2
    assert contents.index("页面说明") < contents.index("[[图片:1]]") < contents.index("[[图片:2]]") < contents.index("结尾")
    assert "右侧说明" in parsed.full_text


def test_Axure原型包_整幅设计稿送图片理解_切图碎片跳过(tmp_path):
    import os
    import zipfile

    import io

    from PIL import Image

    z = tmp_path / "proto.zip"
    big = io.BytesIO(); Image.frombytes("RGB", (900, 700), os.urandom(900 * 700 * 3)).save(big, format="PNG")
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("proto/data/document.js", '$axure.loadDocument({"sitemap":{"rootNodes":[{"pageName":"首页","url":"home.html"}]}});')
        zf.writestr("proto/home.html", "<html><body><div>登录入口</div></body></html>")
        zf.writestr("proto/images/home/u1.png", big.getvalue())          # 整幅设计稿
        zf.writestr("proto/images/home/u2.png", _png_bytes((40, 40)))      # 按钮切图
        zf.writestr("proto/images/home/u3.svg", "<svg/>")                  # 矢量切片
    parsed = parse_file(z)
    assert "Axure 原型" in parsed.full_text and "登录入口" in parsed.full_text
    assert len(parsed.embedded_images) == 1 and "[[图片:1]]" in parsed.full_text


def test_解析内容体检_源码样式判为疑似非需求文本():
    from app.parsers.sanity import content_warning

    css = "\n".join(f"#u{i} {{\n  position:absolute;\n  left:{i}px;\n  width:100px;\n}}" for i in range(200))
    assert "代码 / 样式类行" in (content_warning(css) or "")
    assert content_warning("# 登录需求\n密码错误 5 次锁定 30 分钟。" * 100) is None
    assert content_warning("short {}") is None
    moji = ("# 文件：Φ┐₧Φ╡óµîæµêÿ/data.js\nabcdefghijklmnopqrstuvwxyz 0123456789 " * 60)
    assert "乱码" in (content_warning(moji) or "")


def test_zip内中文文件名没有UTF8标志位_还原后Axure页面名与说明备注不丢(tmp_path):
    """Windows / Axure 打的 zip 条目名常不带 UTF-8 标志位，Python 按 cp437 解码成「µ┤╗σè¿」乱码：
    站点地图里的页面 url 对不上包内文件，页面名退回乱码、每页 data.js 的说明备注全部找不到。"""
    import json
    import zipfile

    class LegacyInfo(zipfile.ZipInfo):
        def _encodeFilenameFlags(self):   # 原始 UTF-8 字节直接落盘，不打 0x800 标志（旧版压缩工具的行为）
            return self.filename.encode("utf-8"), self.flag_bits & ~0x800

    sitemap = {"sitemap": {"rootNodes": [{"pageName": "pc-活动详情页界面说明", "type": "Wireframe", "url": "pc-活动详情页界面说明.html"}]}}
    page = ("<html><head><title>pc-活动详情页界面说明</title></head><body><div id='base'>"
            "<div id='u0' class='ax_default label'><div id='u0_text' class='text'><p><span>报名费 68 元</span></p></div></div></div></body></html>")
    z = tmp_path / "原型.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for name, data in [("data/document.js", "$axure.loadDocument(" + json.dumps(sitemap, ensure_ascii=False) + ");"),
                           ("pc-活动详情页界面说明.html", page),
                           ("files/pc-活动详情页界面说明/data.js",
                            '$axure.loadCurrentPage((function() { var a="<p><span>盈利报名未结清时不能再次报名</span></p>"; return a; })());')]:
            zf.writestr(LegacyInfo(name), data.encode("utf-8"))
    with zipfile.ZipFile(z) as zf:   # 前提成立：Python 原生读出来确实是乱码
        assert any("µ" in i.filename for i in zf.infolist())
    text = parse_file(z).full_text
    assert "页面：pc-活动详情页界面说明" in text and "µ" not in text
    assert "报名费 68 元" in text and "【说明 / 备注】" in text and "盈利报名未结清时不能再次报名" in text


def test_zip条目名GBK编码也能还原(tmp_path):
    import zipfile

    class GbkInfo(zipfile.ZipInfo):
        def _encodeFilenameFlags(self):
            return self.filename.encode("gb18030"), self.flag_bits & ~0x800

    z = tmp_path / "包.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr(GbkInfo("需求/登录需求.md"), "# 登录\n密码错误 5 次锁定 30 分钟".encode("utf-8"))
    text = parse_file(z).full_text
    assert "登录需求.md" in text and "密码错误 5 次锁定 30 分钟" in text


def test_网页转文本_按字体拆开的span不拆行_br处断行():
    from app.parsers.link import html_to_markdown

    html = ("<html><body><div id='base'><div id='u1_text' class='text'>"
            "<p><span>先判断未覆盖基础奖金是否</span><span style='font-family:Helvetica'>≤</span><span>0，是则正常结算</span></p>"
            "<p><span>第一行</span><br><span>第二行</span></p></div></div></body></html>")
    _, body = html_to_markdown(html)
    lines = body.splitlines()
    assert "先判断未覆盖基础奖金是否≤0，是则正常结算" in lines   # 一句话里的运算符不被拆成单独一行
    assert "第一行" in lines and "第二行" in lines


def test_Axure页面文字_按控件整块去重_不按行去重(tmp_path):
    """同页重复的是整个控件（母版 / 面板状态）；两个控件共用的短行（编号、运算符、「各字段详细说明如下」）不能被后一个控件丢掉。"""
    import zipfile

    from app.parsers.base import parse_file

    box = lambda i, *paras: (f"<div id='u{i}' class='ax_default'><div id='u{i}_text' class='text'>"
                             + "".join(f"<p>{p}</p>" for p in paras) + "</div></div>")
    page = ("<html><head><title>需求</title></head><body><div id='base'>"
            + box(1, "<span>2.5.1 字段说明</span>", "<span>各字段详细说明如下</span>")
            + box(2, "<span>2.6.1 字段说明</span>", "<span>各字段详细说明如下</span>")
            + box(3, "<span>剩余应占位金额 </span><span>≤</span><span> 0，则不再需要占位</span>")
            + box(4, "<span>剩余应占位金额 </span><span>≤</span><span> 0，则不再需要占位</span>")  # 面板另一状态整块重复
            + "</div></body></html>")
    z = tmp_path / "原型.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("原型/index.html", "<html><body><iframe src='start.html'></iframe></body></html>")
        zf.writestr("原型/需求.html", page)
        zf.writestr("原型/data/document.js", '$axure.loadDocument({"sitemap":{"rootNodes":[{"pageName":"需求","url":"需求.html"}]}});')
        zf.writestr("原型/resources/scripts/axure/axure.js", "var $axure = {};")
    text = parse_file(z).full_text
    assert text.count("各字段详细说明如下") == 2     # 两个控件各自保留
    assert text.count("剩余应占位金额 ≤ 0，则不再需要占位") == 1   # 整块重复的控件只留一份，且运算符在句子里
    assert "\n≤\n" not in text
    # 每行带控件编号，需求提取的底稿按编号逐控件核对去向；一个控件多段每段都带同一编号
    assert "[u1] 2.5.1 字段说明\n[u1] 各字段详细说明如下\n[u2] 2.6.1 字段说明" in text and "[u3] 剩余应占位金额 ≤ 0" in text and "[u4]" not in text


def test_Axure按坐标还原阅读顺序_表单控件单独抓_隐藏态照提_整页字符自校验(tmp_path):
    """《Axure 需求提取流程》：控件带坐标提取（子控件坐标相对父控件）、按（行 4px 容差，列）排序；下拉框选项 / 输入框值不在 _text
    里要单独抓；隐藏态照样提取并标记；整页 get_text 字符多重集减提取行必须为空（与控件解析相互独立），漏了就报。"""
    import zipfile

    from app.parsers.base import parse_file

    def box(i, text, cls="_段落", extra=""):
        return (f"<div id='u{i}' class='ax_default {cls}'{extra}><div id='u{i}_div'></div>"
                f"<div id='u{i}_text' class='text'><p><span>{text}</span></p></div></div>")
    # DOM 顺序 u1(下方) → u2(上方右) → u3(上方左)；u4 是组合，子控件 u5 / u6 坐标相对 u4；u7 隐藏；u8 下拉框；u9 输入框
    page = ("<html><head><title>流程</title></head><body><div id='base'>"
            + box(1, "结余部分计入第二天") + box(2, "否", "label") + box(3, "是否有结余", "flow_shape")
            + "<div id='u4' class='ax_default'><div id='u4_div'></div>" + box(6, "第二行") + box(5, "第一行") + "</div>"
            + box(7, "隐藏的弹窗文案", "box_1", " style=\"display:none; visibility: hidden\"")
            + "<div id='u8' class='ax_default droplist'><select id='u8_input'><option value='全部'>全部</option><option value='日榜'>日榜</option></select></div>"
            + "<div id='u9' class='ax_default text_field'><input id='u9_input' type='text' value='100'/></div>"
            + "<!-- Unnamed (矩形) --></div></body></html>")
    css = ("#u1 { position:absolute; left:10px; top:300px; width:100px; height:20px; }"
           "#u2 { left:200px; top:102px; } #u3 { left:10px; top:100px; }"
           "#u4 { left:0px; top:400px; } #u5 { left:0px; top:0px; } #u6 { left:0px; top:30px; }"
           "#u7 { left:0px; top:500px; } #u8 { left:0px; top:600px; } #u9 { left:0px; top:700px; }")
    orphan = ("<html><head><title>孤儿</title></head><body><div id='base'>" + box(21, "有编号的行")
              + "<div class='stray'>没有控件编号的文字</div></div></body></html>")
    z = tmp_path / "原型.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("原型/index.html", "<html><body><iframe src='start.html'></iframe></body></html>")
        zf.writestr("原型/流程.html", page)
        zf.writestr("原型/files/流程/styles.css", css)
        zf.writestr("原型/孤儿.html", orphan)
        zf.writestr("原型/data/document.js", '$axure.loadDocument({"sitemap":{"rootNodes":[{"pageName":"流程","url":"流程.html"},{"pageName":"孤儿","url":"孤儿.html"}]}});')
        zf.writestr("原型/resources/scripts/axure/axure.js", "var $axure = {};")
    doc = parse_file(z)
    text = doc.full_text
    # 阅读顺序按坐标：上方左 → 上方右（同一行 4px 容差）→ 下方 → 组合内按相对坐标 → 隐藏控件 → 表单控件
    assert ("[u3] 是否有结余\n[u2] 否\n[u1] 结余部分计入第二天\n[u5] 第一行\n[u6] 第二行\n[u7] 隐藏的弹窗文案\n"
            "[u8] 下拉选项：全部 / 日榜\n[u9] 输入值：100") in text
    assert doc.controls["u5"] == {"cls": "_段落", "x": 0, "y": 400} and doc.controls["u6"]["y"] == 430   # 子控件绝对坐标沿父链累加
    assert doc.controls["u7"] == {"cls": "box_1", "x": 0, "y": 500, "hidden": True} and doc.controls["u3"]["cls"] == "flow_shape"
    assert doc.controls["u8"]["cls"] == "droplist" and "u4" not in doc.controls   # 没有文字 / 图片的容器不进索引
    pages = {p["page"]: p for p in doc.check["pages"]}
    ok = pages["流程"]
    assert ok["missing"] == 0 and ok["raw"] == ok["got"] > 0 and ok["controls"] == 8 and ok["hidden"] == 1 and "outside" not in ok
    bad = pages["孤儿"]
    assert bad["missing"] == len("没有控件编号的文字") and "没" in bad["missing_chars"] and bad["outside"] == ["没有控件编号的文字"]
    assert doc.check["ok"] is False and doc.check["missing"] == bad["missing"]
    assert "没有控件编号的文字" not in text   # 自校验就是为了把这种漏抓暴露出来，而不是静默过


def test_Axure表格同一行单元格并成一行_表头与数据行对应_自校验不计分隔符(tmp_path):
    """表格单元格各是一个控件，按行各自成行会丢掉「条件 → 结果」的对应（「输赢金额=0」与「不变」分家）。同一行（top 相同）的
    单元格并成「A ｜ B」一条，编号取第一格，cls=table_row；表头行也是一行；跨行合并的单元格单独一行；整页字符自校验不计「｜」。"""
    import zipfile

    from app.parsers.base import parse_file

    def cell(i, text):
        return (f"<div id='u{i}' class='ax_default table_cell'><img id='u{i}_img' src='images/u{i}.png'/>"
                f"<div id='u{i}_text' class='text'><p><span>{text}</span></p></div></div>")
    page = ("<html><head><title>表</title></head><body><div id='base'><div id='u10' class='ax_default'>"
            + cell(14, "当前连赢次数“不变”") + cell(13, "输赢金额=0") + cell(12, "系统计算方式") + cell(11, "本局游戏结果") + cell(15, "合并单元格说明")
            + "</div></div></body></html>")
    css = ("#u10 { left:100px; top:200px; } #u11 { left:0px; top:0px; } #u12 { left:200px; top:1px; }"
           "#u13 { left:0px; top:30px; } #u14 { left:200px; top:30px; } #u15 { left:0px; top:60px; }")
    z = tmp_path / "表.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("原型/index.html", "<html><body><iframe src='start.html'></iframe></body></html>")
        zf.writestr("原型/表.html", page)
        zf.writestr("原型/files/表/styles.css", css)
        zf.writestr("原型/data/document.js", '$axure.loadDocument({"sitemap":{"rootNodes":[{"pageName":"表","url":"表.html"}]}});')
        zf.writestr("原型/resources/scripts/axure/axure.js", "var $axure = {};")
    doc = parse_file(z)
    assert "[u11] 本局游戏结果 ｜ 系统计算方式\n[u13] 输赢金额=0 ｜ 当前连赢次数“不变”\n[u15] 合并单元格说明" in doc.full_text
    assert doc.controls["u11"] == {"cls": "table_row", "x": 100, "y": 200, "img": "images/u11.png"} and doc.controls["u15"]["cls"] == "table_cell"
    assert "u12" not in doc.controls   # 并进行里的单元格不单独建索引
    page_check = doc.check["pages"][0]
    assert page_check["missing"] == 0 and page_check["extra"] == 0 and page_check["raw"] == page_check["got"]


def test_Axure画布独立区域_与主体相距千像素的控件群单独成块并标疑似草稿(tmp_path):
    """作者把旧版本丢在画布右侧 / 下方忘了删：与主体任何控件横纵都相距 ≥ 1000px 的控件群（≥ 3 个）算独立区域，主体先输出、
    区域各自成块、前面一行【独立区域 N：…疑似遗留草稿】标记（≤ 40 字，底稿不当原文句），控件索引记 region；零星 1～2 个不算。"""
    import zipfile

    from app.parsers.base import parse_file

    def box(i, text):
        return f"<div id='u{i}' class='ax_default _段落'><div id='u{i}_text' class='text'><p>{text}</p></div></div>"
    page = ("<html><head><title>记录</title></head><body><div id='base'>" + box(1, "正式页面标题") + box(2, "日榜说明") + box(3, "月榜说明")
            + box(11, "旧版：年榜说明") + box(12, "晋级年榜") + box(13, "道具明细") + box(21, "远处零星一个备注") + "</div></body></html>")
    css = ("#u1 { left:30px; top:40px; } #u2 { left:30px; top:100px; } #u3 { left:30px; top:160px; }"
           "#u11 { left:2400px; top:40px; } #u12 { left:2400px; top:100px; } #u13 { left:2420px; top:160px; } #u21 { left:30px; top:5000px; }")
    z = tmp_path / "记录.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("原型/index.html", "<html><body><iframe src='start.html'></iframe></body></html>")
        zf.writestr("原型/记录.html", page)
        zf.writestr("原型/files/记录/styles.css", css)
        zf.writestr("原型/data/document.js", '$axure.loadDocument({"sitemap":{"rootNodes":[{"pageName":"记录","url":"记录.html"}]}});')
        zf.writestr("原型/resources/scripts/axure/axure.js", "var $axure = {};")
    doc = parse_file(z)
    assert ("[u1] 正式页面标题\n[u2] 日榜说明\n[u3] 月榜说明\n[u21] 远处零星一个备注\n"
            "【独立区域 1：主体右侧 2370px，3 控件，疑似遗留草稿】\n[u11] 旧版：年榜说明\n[u12] 晋级年榜\n[u13] 道具明细") in doc.full_text
    assert doc.controls["u11"]["region"] == 1 and "region" not in doc.controls["u1"] and "region" not in doc.controls["u21"]
    assert doc.check["pages"][0]["islands"] == 1 and doc.check["ok"] is True
    from app.agents.ledger import build_units
    units = build_units("【文件：记录.zip】\n# Axure 原型：记录.zip\n## 页面：记录\n" + doc.full_text.split("## 页面：记录\n", 1)[-1])
    assert all("独立区域" not in u["text"] for u in units) and any(u["uid"] == "u11" for u in units)   # 标记行不进底稿
