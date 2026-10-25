"""生成质量核心闭环测试（AI测试用例生成工具核心需求设计）：

测试点实体与审核状态机、过粗检测、查漏补缺（只新增）、重复检查、
驳回定点修改（锁定保护 + Diff）、需求变更最小范围更新、学习候选与规则库。
"""

import json

import httpx
import pytest
from asgi_lifespan import LifespanManager

from app.main import app
from app.tasks.points import (
    add_points,
    apply_point_review,
    assign_entities,
    case_duplicate_candidates,
    coarse_warnings,
    confirmable_points,
    duplicate_candidates,
)
from tests.stubs import StubLLM, generator_reply, make_case, review_reply


@pytest.fixture
async def client():
    async with LifespanManager(app):
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
            yield c


def _modules():
    return assign_entities([
        {"module": "登录", "points": [
            {"point": "正确账号密码登录成功", "dimension": "正常流程"},
            {"point": "连续错误5次账号锁定", "dimension": "异常流程"},
        ]},
    ])


ANALYST_REPLY_V2 = json.dumps({
    "modules": [{"module": "登录", "points": [
        {"point": "正确账号密码登录成功", "dimension": "正常流程"},
        {"point": "连续错误5次账号锁定", "dimension": "异常流程"},
    ]}],
    "blind_spots": [],
}, ensure_ascii=False)

GAP_REPLY = json.dumps({
    "coverage": {"正常流程": "已覆盖", "异常流程": "已覆盖", "输入校验": "未覆盖",
                 "边界值": "已覆盖", "状态转换": "不适用", "权限": "不适用", "数据": "已覆盖",
                 "重复操作": "未覆盖", "中断": "不适用", "网络": "不适用"},
    "additions": [{"module": "登录", "point": "密码为空点击登录被拦截", "dimension": "输入校验",
                   "reason": "输入校验未覆盖"}],
}, ensure_ascii=False)


# ---- 测试点实体化与状态机 ----


def test_assign_entities_幂等与tp_id():
    modules = _modules()
    tp_ids = [p["tp_id"] for p in modules[0]["points"]]
    assert tp_ids == ["TP001", "TP002"]
    assert all(p["status"] == "pending" and not p["locked"] for p in modules[0]["points"])
    again = assign_entities(modules)  # 幂等：不重新分配 tp_id、不丢状态
    assert [p["tp_id"] for p in again[0]["points"]] == tp_ids


def test_coarse_detection_过粗测试点提示():
    assert coarse_warnings("验证登录整体功能是否正常")
    assert coarse_warnings("覆盖所有情况")
    assert not coarse_warnings("连续错误5次账号锁定")


def test_point_review_approve_即锁定():
    modules = _modules()
    out = apply_point_review(modules, [{"tp_id": "TP001", "action": "approve"}])
    p = modules[0]["points"][0]
    assert (p["status"], p["locked"]) == ("approved", True)
    assert out["counts"]["approve"] == 1


def test_point_review_reject_须带意见且连续驳回提示():
    modules = _modules()
    with pytest.raises(ValueError):
        apply_point_review(modules, [{"tp_id": "TP001", "action": "reject"}])
    with pytest.raises(ValueError):  # 完整需求 6.4：驳回类型必填
        apply_point_review(modules, [{"tp_id": "TP001", "action": "reject", "comment": "过粗"}])
    apply_point_review(modules, [{"tp_id": "TP001", "action": "reject", "comment": "颗粒度过粗",
                                  "reject_types": ["颗粒度过粗"]}])
    out = apply_point_review(modules, [{"tp_id": "TP001", "action": "reject", "comment": "仍然过粗",
                                        "reject_types": ["颗粒度过粗"]}])
    assert modules[0]["points"][0]["reject_count"] == 2
    assert out["hints"] and "TP001" in out["hints"][0]  # 需求三十五：连续驳回提示人工介入


def test_confirmable_points_只取通过项且驳回项永不进入():
    modules = _modules()
    # 全部未审核：沿用全量
    assert sum(len(e["points"]) for e in confirmable_points(modules)) == 2
    apply_point_review(modules, [
        {"tp_id": "TP001", "action": "approve"},
        {"tp_id": "TP002", "action": "reject", "comment": "重复", "reject_types": ["重复"]},
    ])
    confirmed = confirmable_points(modules)
    assert [p["point"] for e in confirmed for p in e["points"]] == ["正确账号密码登录成功"]


def test_add_points_只新增且生成前查重():
    modules = _modules()
    added = add_points(modules, [
        {"module": "登录", "point": "密码为空点击登录被拦截", "dimension": "输入校验"},
        {"module": "登录", "point": "正确账号密码登录成功"},  # 与已有高度相似：跳过
    ], source="gap")
    assert len(added) == 1 and added[0]["source"] == "gap" and added[0]["status"] == "pending"
    assert len(modules[0]["points"]) == 3  # 存量 2 条不动


def test_duplicate_candidates_文字初筛():
    modules = assign_entities([{"module": "登录", "points": [
        "验证密码为空无法登录", "验证密码为空时登录失败", "连续错误5次账号锁定",
    ]}])
    pairs = duplicate_candidates(modules)
    assert len(pairs) == 1
    assert {pairs[0]["a"], pairs[0]["b"]} == {"TP001", "TP002"}


def test_case_duplicate_candidates_综合字段比对():
    a = make_case()
    b = make_case(case_id="TC-登录-002")  # 同标题同步骤
    c = make_case(case_id="TC-登录-003", title="验证锁定策略",
                  steps=[{"action": "连续输错5次密码", "expected": "账号被锁定"}])
    pairs = case_duplicate_candidates([a, b, c])
    assert len(pairs) == 1 and {pairs[0]["a"], pairs[0]["b"]} == {"TC-登录-001", "TC-登录-002"}


# ---- 定点修改确定性合并 ----


def test_apply_point_fixes_split与越权保护():
    from app.agents.quality import apply_point_fixes

    modules = _modules()
    apply_point_review(modules, [{"tp_id": "TP002", "action": "reject", "comment": "颗粒度过粗",
                                  "reject_types": ["颗粒度过粗"]}])
    data = {"fixes": [
        {"tp_id": "TP002", "comment_type": "颗粒度过粗", "action": "split",
         "split_into": [{"point": "密码错误登录失败", "dimension": "异常流程"},
                        {"point": "连续错误5次账号锁定", "dimension": "异常流程"}]},
        {"tp_id": "TP001", "action": "modify", "point": "越权改动"},  # 非驳回项：忽略
    ]}
    out = apply_point_fixes(modules, data, allowed={"TP002"})
    points = modules[0]["points"]
    assert [p["tp_id"] for p in points] == ["TP001", "TP002-01", "TP002-02"]
    assert points[0]["point"] == "正确账号密码登录成功"  # 越权改动被丢弃
    assert all(p["status"] == "pending" for p in points[1:])  # 局部修改 → 再审核
    assert len(out["diff"]) == 1 and out["diff"][0]["action"] == "split"


def test_merge_case_fix_锁定保护与顺延编号():
    from app.agents.quality import merge_case_fix

    cases = [
        dict(make_case(), uid="u1"),
        dict(make_case(case_id="TC-登录-002", title="验证密码错误提示"), uid="u2"),
    ]
    reviews = {
        "u1": {"status": "approved", "locked": True, "comment": "", "reject_count": 0},
        "u2": {"status": "rejected", "locked": False, "comment": "预期不可验证", "reject_count": 1},
    }
    data = {
        "cases": [
            dict(make_case(), title="越权修改锁定用例"),                      # TC-登录-001 锁定：丢弃
            dict(make_case(case_id="TC-登录-002", title="验证密码错误提示文案")),  # 允许修改
            dict(make_case(case_id="TC-登录-009", title="新增：锁定后正确密码仍拒绝")),  # 新增
        ],
        "deleted": [],
    }
    out = merge_case_fix(cases, data, allowed={"TC-登录-002"}, reviews=reviews)
    result = out["cases"]
    assert result[0]["title"] == "验证正确账号密码登录成功"  # 锁定用例原样
    assert result[1]["title"] == "验证密码错误提示文案"
    assert result[2]["case_id"] == "TC-登录-003"  # 新增顺延编号，不重排存量
    assert reviews[result[1]["uid"]]["status"] == "pending"  # 修改后回到待审核
    actions = {d["action"] for d in out["diff"]}
    assert actions == {"modify", "add"}


# ---- 规则库（需求三十六~三十九）----


def test_rule_store_候选确认与范围过滤(tmp_path):
    from app.learning import RuleStore

    store = RuleStore(tmp_path / "rules.json")
    added = store.add_candidates([
        {"content": "支付场景默认加入快速重复点击", "evidence": "12条中10条人工增加", "occurrences": 10,
         "confidence": "高", "scope_hint": "project"},
        {"content": "支付场景默认加入快速重复点击", "occurrences": 3},  # 重复：去重
    ], project="德州扑克")
    assert len(added) == 1 and added[0].status == "candidate"
    # 候选不生效：active_rules 为空
    assert store.active_rules("德州扑克") == []
    with pytest.raises(ValueError):
        store.confirm(added[0].rule_id, "project", project=None)  # 项目级必须给项目名
    store.confirm(added[0].rule_id, "project", project="德州扑克")
    assert len(store.active_rules("德州扑克")) == 1
    assert store.active_rules("斗地主") == []  # 项目隔离（需求三十九）
    notes, snapshot = store.render("德州扑克")
    assert "快速重复点击" in notes and snapshot[0]["scope"] == "project"


# ---- API：测试点审核工作台 ----


async def _create_confirm_task(client, extra_replies=()):
    app.state.llm = StubLLM([ANALYST_REPLY_V2, GAP_REPLY, *extra_replies])
    resp = await client.post(
        "/api/v1/tasks",
        data={"text": "登录需求：账号密码登录，连续错误5次锁定", "confirm_points": "true"},
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


async def test_confirm_flow_实体化与自动查漏(client):
    data = await _create_confirm_task(client)
    # 拆解 2 条 + 查漏新增 1 条（只新增，存量不动）
    points = data["test_points"][0]["points"]
    assert [p["tp_id"] for p in points] == ["TP001", "TP002", "TP003"]
    assert points[2]["source"] == "gap"
    assert data["coverage"]["输入校验"] == "未覆盖"


async def test_points_review_批量与驳回后定点修改(client):
    data = await _create_confirm_task(client)
    task_id = data["task_id"]
    # 批量通过 + 驳回（需求四十九/五十）
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [
        {"tp_id": "TP001", "action": "approve"},
        {"tp_id": "TP003", "action": "approve"},
        {"tp_id": "TP002", "action": "reject", "comment": "测试点过大，需要拆分登录失败和账号锁定",
         "reject_types": ["颗粒度过粗"]},
    ]})
    assert resp.status_code == 200
    assert resp.json()["counts"] == {"approve": 2, "reject": 1, "modify": 0, "delete": 0, "unlock": 0}

    # AI 定点修改：只处理被驳回的 TP002，拆分为两条（需求三十三）
    app.state.llm = StubLLM([json.dumps({"fixes": [
        {"tp_id": "TP002", "comment_type": "颗粒度过粗", "action": "split",
         "split_into": [{"point": "密码错误登录失败", "dimension": "异常流程"},
                        {"point": "连续错误5次账号锁定", "dimension": "异常流程"}]},
    ], "additions": []}, ensure_ascii=False)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/fix")
    assert resp.status_code == 200
    pending = resp.json()["pending_fix"]
    assert pending["kind"] == "points" and len(pending["proposals"]) == 1
    assert pending["proposals"][0]["action"] == "split"
    # 确认流（完整需求 7.3）：接受后才落地
    resp = await client.post(f"/api/v1/tasks/{task_id}/fix/confirm", json={"accept_all": True})
    assert resp.status_code == 200
    fixed = resp.json()
    assert fixed["diff"][0]["comment_type"] == "颗粒度过粗"
    tp_ids = [p["tp_id"] for p in fixed["test_points"][0]["points"]]
    assert "TP002-01" in tp_ids and "TP002-02" in tp_ids

    # 确认生成：只有已通过的点进入正式测试点（TP002-xx 尚未通过）
    app.state.llm = StubLLM([generator_reply(make_case()), review_reply(True)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})
    assert resp.status_code == 200
    gen_prompt = app.state.llm.calls[0]["messages"][1]["content"]
    assert "正确账号密码登录成功" in gen_prompt
    assert "密码错误登录失败" not in gen_prompt  # 未通过的点不进入生成


async def test_points_add_ai补充只新增(client):
    data = await _create_confirm_task(client)
    task_id = data["task_id"]
    app.state.llm = StubLLM([json.dumps({"additions": [
        {"module": "登录", "point": "网络切换后重新登录保持会话", "dimension": "网络"},
    ]}, ensure_ascii=False)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/add",
                             json={"instruction": "补充网络切换场景"})
    assert resp.status_code == 200
    added = resp.json()["added"]
    assert len(added) == 1 and added[0]["source"] == "supplement"


# ---- API：用例驳回定点修改与锁定 ----


async def _create_completed_task(client, *cases):
    from tests.stubs import ANALYST_REPLY

    app.state.llm = StubLLM([ANALYST_REPLY, generator_reply(*cases), review_reply(True)])
    resp = await client.post("/api/v1/tasks", data={"text": "登录需求"})
    return resp.json()["task_id"]


async def test_case_reject_then_fix_锁定保护(client):
    task_id = await _create_completed_task(
        client, make_case(), make_case(case_id="TC-登录-002", title="验证密码错误提示"),
    )
    resp = await client.post(f"/api/v1/tasks/{task_id}/review", json={"items": [
        {"case_id": "TC-登录-001", "action": "approve"},
        {"case_id": "TC-登录-002", "action": "reject", "comment": "预期结果空泛不可验证",
         "reject_types": ["预期不可验证"]},
    ]})
    assert resp.status_code == 200

    fixed_case = make_case(case_id="TC-登录-002", title="验证密码错误提示",
                           steps=[{"action": "输入错误密码提交", "expected": "提示「密码错误，还可尝试4次」"}])
    app.state.llm = StubLLM([json.dumps({
        "fixes": [{"case_id": "TC-登录-002", "comment_type": "预期不可验证", "note": "补充具体文案"}],
        "cases": [fixed_case], "deleted": [],
    }, ensure_ascii=False)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/cases/fix")
    assert resp.status_code == 200
    pending = resp.json()["pending_fix"]
    assert pending["proposals"][0]["comment_type"] == "预期不可验证"
    # 修改仅发送被驳回用例（需求三十一）
    sent = app.state.llm.calls[0]["messages"][1]["content"]
    assert "TC-登录-002" in sent and "TC-登录-001" not in sent
    # 确认流（完整需求 9.4）：接受后才落地
    resp = await client.post(f"/api/v1/tasks/{task_id}/fix/confirm", json={"accept_all": True})
    assert resp.status_code == 200

    task = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    titles = {c["case_id"]: c for c in task["result"]["cases"]}
    assert "还可尝试4次" in titles["TC-登录-002"]["steps"][0]["expected"]
    # 锁定用例原样保留，状态机：修改后的用例回到 pending
    states = list(task["case_reviews"].values())
    assert sorted(s["status"] for s in states) == ["approved", "pending"]


async def test_revise_不动锁定用例(client):
    task_id = await _create_completed_task(
        client, make_case(), make_case(case_id="TC-登录-002", title="验证密码错误提示"),
    )
    await client.post(f"/api/v1/tasks/{task_id}/review", json={"items": [
        {"case_id": "TC-登录-001", "action": "approve"},
    ]})
    app.state.llm = StubLLM([
        json.dumps({"cases": [dict(make_case(case_id="TC-登录-001", title="修订后的密码错误提示"))],
                    "deleted": []}, ensure_ascii=False),
        review_reply(True),
    ])
    resp = await client.post(f"/api/v1/tasks/{task_id}/revise",
                             json={"instruction": "密码错误提示补充剩余次数"})
    assert resp.status_code == 200
    # 修订上下文只含未锁定用例
    sent = app.state.llm.calls[0]["messages"][1]["content"]
    assert "验证密码错误提示" in sent and "验证正确账号密码登录成功" not in sent
    task = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    titles = [c["title"] for c in task["result"]["cases"]]
    assert "验证正确账号密码登录成功" in titles  # 锁定用例完好


async def test_confirm_points_后台执行可见可管理(client):
    """先审核测试点 + 后台执行：任务立即落库可见，拆解完成后进入待确认。"""
    import asyncio

    app.state.llm = StubLLM([ANALYST_REPLY_V2, GAP_REPLY])
    resp = await client.post(
        "/api/v1/tasks",
        data={"text": "登录需求", "confirm_points": "true", "async_mode": "true"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "queued"  # 立即返回，任务已可在列表查看
    task_id = body["task_id"]

    for _ in range(100):  # 等待后台拆解 + 查漏完成
        await asyncio.sleep(0.02)
        task = (await client.get(f"/api/v1/tasks/{task_id}")).json()
        if task["status"] == "awaiting_confirmation" and task["progress"] is None:
            break
    assert task["status"] == "awaiting_confirmation"
    points = task["analysis"]["test_points"][0]["points"]
    assert points[0]["tp_id"] == "TP001"  # 实体化完成
    assert task["coverage"]["输入校验"] == "未覆盖"  # 后台查漏已补充


async def test_cancel_运行中的后台任务(client):
    """任务管理：进行中的后台任务可取消，标记失败并留痕。"""
    import asyncio

    class HangLLM:  # 模拟长时间运行的模型调用
        async def chat(self, *args, **kwargs):
            await asyncio.Event().wait()

    app.state.llm = HangLLM()
    resp = await client.post("/api/v1/tasks", data={"text": "登录需求", "async_mode": "true"})
    task_id = resp.json()["task_id"]
    await asyncio.sleep(0.05)  # 等后台任务启动

    resp = await client.post(f"/api/v1/tasks/{task_id}/cancel")
    assert resp.status_code == 200 and resp.json()["canceled"] is True
    task = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert task["status"] == "failed" and "取消" in task["error"]
    # 重复取消：无可取消的执行
    assert (await client.post(f"/api/v1/tasks/{task_id}/cancel")).status_code == 409


async def test_task_list_项目名展示与过滤(client):
    from tests.stubs import ANALYST_REPLY

    app.state.llm = StubLLM([ANALYST_REPLY, generator_reply(make_case()), review_reply(True)])
    resp = await client.post("/api/v1/tasks", data={"text": "登录需求", "project": "德州扑克"})
    assert resp.status_code == 200

    data = (await client.get("/api/v1/tasks")).json()
    assert any(t["project"] == "德州扑克" for t in data["tasks"])
    filtered = (await client.get("/api/v1/tasks", params={"project": "德州扑克"})).json()
    assert filtered["tasks"] and all(t["project"] == "德州扑克" for t in filtered["tasks"])
    empty = (await client.get("/api/v1/tasks", params={"project": "不存在的项目"})).json()
    assert empty["tasks"] == []


async def test_项目增删改查与改名联动(client):
    # 新建 + 重名拒绝
    resp = await client.post("/api/v1/projects", json={"name": "增删改查", "description": "CRUD 验证"})
    assert resp.status_code == 200
    assert (await client.post("/api/v1/projects", json={"name": "增删改查"})).status_code == 400

    # 创建任务时选择该项目（自动注册路径同样兼容新名称）
    from tests.stubs import ANALYST_REPLY

    app.state.llm = StubLLM([ANALYST_REPLY, generator_reply(make_case()), review_reply(True)])
    resp = await client.post("/api/v1/tasks", data={"text": "登录需求", "project": "增删改查"})
    task_id = resp.json()["task_id"]

    # 列表带描述与统计
    projects = (await client.get("/api/v1/projects")).json()["projects"]
    row = next(p for p in projects if p["project"] == "增删改查")
    assert row["description"] == "CRUD 验证" and row["tasks"] == 1

    # 有任务不可删；改名联动任务归属
    assert (await client.delete("/api/v1/projects/增删改查")).status_code == 400
    resp = await client.put("/api/v1/projects/增删改查",
                            json={"name": "增删改查V2", "description": "已改名"})
    assert resp.status_code == 200
    task = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert task["context"]["project"] == "增删改查V2"
    listed = (await client.get("/api/v1/tasks", params={"project": "增删改查V2"})).json()["tasks"]
    assert any(t["task_id"] == task_id for t in listed)

    # 空项目可删
    await client.post("/api/v1/projects", json={"name": "空项目"})
    assert (await client.delete("/api/v1/projects/空项目")).status_code == 200
    assert (await client.delete("/api/v1/projects/空项目")).status_code == 404


async def test_用例执行轮次与执行记录(client):
    task_id = await _create_completed_task(
        client, make_case(), make_case(case_id="TC-登录-002", title="验证密码错误提示"),
    )
    # 新建轮次
    run = (await client.post(f"/api/v1/tasks/{task_id}/executions", json={"name": "冒烟"})).json()
    assert run["name"] == "冒烟" and run["summary"]["total"] == 2
    # 未结束前不可再开新轮次
    assert (await client.post(f"/api/v1/tasks/{task_id}/executions", json={})).status_code == 409

    rid = run["run_id"]
    # 失败必须填原因；未知状态拒绝
    assert (await client.post(f"/api/v1/tasks/{task_id}/executions/{rid}/results",
            json={"items": [{"case_id": "TC-登录-001", "status": "fail"}]})).status_code == 400
    assert (await client.post(f"/api/v1/tasks/{task_id}/executions/{rid}/results",
            json={"items": [{"case_id": "TC-登录-001", "status": "ok"}]})).status_code == 400
    # 批量记录：1 通过 + 1 失败（带缺陷号）
    resp = await client.post(f"/api/v1/tasks/{task_id}/executions/{rid}/results", json={"items": [
        {"case_id": "TC-登录-001", "status": "pass"},
        {"case_id": "TC-登录-002", "status": "fail", "note": "BUG-1024 提示文案错误"},
    ]})
    assert resp.status_code == 200
    summary = resp.json()["summary"]
    assert summary["pass"] == 1 and summary["fail"] == 1 and summary["pass_rate"] == 0.5

    # 复测覆盖并保留历史
    resp = await client.post(f"/api/v1/tasks/{task_id}/executions/{rid}/results", json={"items": [
        {"case_id": "TC-登录-002", "status": "pass"},
    ]})
    results = resp.json()["results"]
    retested = next(r for r in results.values() if r["case_id"] == "TC-登录-002")
    assert retested["status"] == "pass" and retested["history"][0]["status"] == "fail"
    assert resp.json()["summary"]["pass_rate"] == 1.0

    # 结束轮次后不可继续记录；可开第二轮（回归）
    assert (await client.post(f"/api/v1/tasks/{task_id}/executions/{rid}/finish")).status_code == 200
    assert (await client.post(f"/api/v1/tasks/{task_id}/executions/{rid}/results",
            json={"items": [{"case_id": "TC-登录-001", "status": "pass"}]})).status_code == 409
    run2 = (await client.post(f"/api/v1/tasks/{task_id}/executions", json={})).json()
    assert run2["name"] == "第 2 轮执行"
    task = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert len(task["executions"]) == 2 and task["executions"][0]["finished_at"]


# ---- API：需求变更最小范围更新 ----


async def test_requirement_diff_analyze_and_apply(client):
    task_id = await _create_completed_task(
        client, make_case(), make_case(case_id="TC-登录-002", title="验证错误3次锁定"),
    )
    app.state.llm = StubLLM([json.dumps({
        "changes": [{"type": "规则调整", "description": "锁定次数 3 次改为 5 次",
                     "affected_points": [], "affected_cases": ["TC-登录-002"],
                     "action_hint": "更新锁定边界"}],
        "new_requirements": [], "summary": "锁定次数调整",
    }, ensure_ascii=False)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/requirement-diff",
                             json={"new_requirement": "账号密码登录，连续错误5次锁定"})
    assert resp.status_code == 200
    assert resp.json()["changes"][0]["affected_cases"] == ["TC-登录-002"]

    updated = make_case(case_id="TC-登录-002", title="验证错误5次锁定")
    app.state.llm = StubLLM([json.dumps({
        "fixes": [], "cases": [updated], "deleted": [],
    }, ensure_ascii=False)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/requirement-diff/apply")
    assert resp.status_code == 200
    task = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    titles = [c["title"] for c in task["result"]["cases"]]
    assert "验证错误5次锁定" in titles and "验证正确账号密码登录成功" in titles
    assert task["context"]["requirement"] == "账号密码登录，连续错误5次锁定"  # 基线更新
    # 重复应用被拒绝
    resp = await client.post(f"/api/v1/tasks/{task_id}/requirement-diff/apply")
    assert resp.status_code == 409


# ---- API：学习候选 → 确认 → 注入生成 ----


async def test_learning_loop_候选确认与注入(client):
    task_id = await _create_completed_task(client, make_case())
    # 制造人工修改留痕
    await client.post(f"/api/v1/tasks/{task_id}/review", json={"items": [
        {"case_id": "TC-登录-001", "action": "modify", "feedback": "统一步骤精简",
         "case": make_case(steps=[{"action": "登录并校验", "expected": "进入首页"}])},
    ]})
    app.state.llm = StubLLM([json.dumps({"candidates": [
        {"content": "登录场景步骤保持精简，单条用例不超过3步", "evidence": "人工修改 2 次",
         "occurrences": 2, "confidence": "中", "scope_hint": "team"},
    ]}, ensure_ascii=False)])
    resp = await client.post("/api/v1/learning/analyze", json={})
    assert resp.status_code == 200
    candidates = resp.json()["candidates"]
    assert len(candidates) == 1

    rule_id = candidates[0]["rule_id"]
    resp = await client.post(f"/api/v1/learning/rules/{rule_id}/confirm", json={"scope": "team"})
    assert resp.status_code == 200 and resp.json()["status"] == "active"

    # 新任务生成时注入规则（需求六十一：反哺下一次生成）
    from tests.stubs import ANALYST_REPLY

    app.state.llm = StubLLM([ANALYST_REPLY, generator_reply(make_case()), review_reply(True)])
    resp = await client.post("/api/v1/tasks", data={"text": "登录需求"})
    assert resp.status_code == 200
    gen_prompt = app.state.llm.calls[1]["messages"][1]["content"]
    assert "步骤保持精简" in gen_prompt and "团队/项目测试规则" in gen_prompt


async def test_后台任务先入库再解析_文件解析失败留痕(client):
    """上传后立即可见（进度 parsing），解析失败的任务标为失败而不是消失；解析成功进入拆解。"""
    import asyncio

    # 不支持的格式：解析在后台失败，任务留痕
    app.state.llm = StubLLM([])
    resp = await client.post("/api/v1/tasks", files={"files": ("需求.xyz", b"\x00\x01\xff\xfe", "application/octet-stream")},
                             data={"confirm_points": "true", "async_mode": "true"})
    assert resp.status_code == 200 and resp.json()["status"] == "queued"
    task_id = resp.json()["task_id"]
    listed = (await client.get("/api/v1/tasks")).json()["tasks"]
    assert any(t["task_id"] == task_id and t["sources"] == ["需求.xyz"] for t in listed)   # 解析前已在列表
    for _ in range(100):
        await asyncio.sleep(0.02)
        task = (await client.get(f"/api/v1/tasks/{task_id}")).json()
        if task["status"] == "failed":
            break
    assert task["status"] == "failed" and task["error"] and task["progress"] is None
    assert (await client.post(f"/api/v1/tasks/{task_id}/retry")).status_code == 409   # 无需求文本不可重试
    # 正常文件：后台解析 → 拆解 → 待确认
    app.state.llm = StubLLM([ANALYST_REPLY_V2, GAP_REPLY])
    resp = await client.post("/api/v1/tasks", files={"files": ("需求.txt", "登录需求：账号密码登录".encode(), "text/plain")},
                             data={"confirm_points": "true", "async_mode": "true"})
    task_id = resp.json()["task_id"]
    for _ in range(100):
        await asyncio.sleep(0.02)
        task = (await client.get(f"/api/v1/tasks/{task_id}")).json()
        if task["status"] == "awaiting_confirmation" and task["progress"] is None:
            break
    assert task["status"] == "awaiting_confirmation" and task["sources"] == ["需求.txt"]
    assert "账号密码登录" in task["context"]["requirement"]


def test_字段级定位_操作步骤含预期_越界改动给出原因():
    from app.agents.quality import build_case_proposals
    origin = make_case(steps=[{"action": "输入密码", "expected": "掩码显示"}, {"action": "点击登录", "expected": "跳转首页"}])
    origin["uid"] = "u1"
    raw = {**origin, "steps": [{"action": "输入密码", "expected": "密码框以圆点掩码显示 8 位"}, {"action": "点击登录", "expected": "跳转首页并右上角展示昵称"}]}
    data = {"fixes": [{"case_id": origin["case_id"], "comment_type": "预期不可验证"}], "cases": [raw], "deleted": []}
    # 勾「操作步骤」+ 第 1 步：第 1 步的预期改动应被采纳，第 2 步还原
    reviews = {"u1": {"status": "rejected", "fields": ["steps"], "steps": [1]}}
    out = build_case_proposals([origin], data, {origin["case_id"]}, reviews)
    assert len(out["proposals"]) == 1
    steps = out["proposals"][0]["after"]["steps"]
    assert steps[0]["expected"] == "密码框以圆点掩码显示 8 位" and steps[1]["expected"] == "跳转首页"
    # 只允许改标题时，模型只改了步骤：不再静默无提案，而是给出原因
    reviews = {"u1": {"status": "rejected", "fields": ["title"], "steps": []}}
    out = build_case_proposals([origin], data, {origin["case_id"]}, reviews)
    assert out["proposals"] == [] and "不在驳回时指定的范围内" in out["invalid"][0]["problem"]


# ---- 人工编写测试点 vs AI 生成测试点对比 ----


def test_parse_manual_points_多种写法():
    from app.tasks.points import parse_manual_points

    text = "1. 登录|密码为空提示\n- 找回密码：验证码过期\nTP03、密码错误5次锁定\n\n   "
    assert parse_manual_points(text, "账号") == [
        {"module": "登录", "point": "密码为空提示"},
        {"module": "找回密码", "point": "验证码过期"},
        {"module": "账号", "point": "密码错误5次锁定"},
    ]


def test_add_points_人工录入不做相似跳过_并给出人工与AI候选配对():
    from app.tasks.points import add_points, assign_entities, cross_source_candidates

    modules = assign_entities(_modules())
    ai_point = modules[0]["points"][0]["point"]
    skipped: list = []
    # 默认相似跳过：与 AI 相同的人工点会被丢掉
    assert add_points(modules, [{"module": "登录", "point": ai_point}], source="manual", skipped=skipped) == []
    assert skipped
    added = add_points(modules, [{"module": "登录", "point": ai_point}, {"module": "登录", "point": "锁定期间不允许找回密码"}],
                       source="manual", skip_similar=False)
    assert len(added) == 2 and all(p["source"] == "manual" for p in added)
    pairs = cross_source_candidates(modules)
    assert pairs and pairs[0]["manual"] == added[0]["tp_id"] and pairs[0]["similarity"] == 1.0


async def test_points_compare_人工与AI逐条匹配与覆盖率(client):
    data = await _create_confirm_task(client)
    task_id = data["task_id"]
    ai_ids = [p["tp_id"] for m in data["test_points"] for p in m["points"]]
    # 没有人工测试点时不能对比
    assert (await client.post(f"/api/v1/tasks/{task_id}/points/compare")).status_code == 409
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/manual-batch",
                             json={"text": "登录|正确账号密码登录成功\n登录|锁定期间不允许找回密码\n", "module": ""})
    assert resp.status_code == 200, resp.text
    manual = resp.json()["added"]
    assert len(manual) == 2 and all(p["source"] == "manual" for p in manual)
    assert (await client.post(f"/api/v1/tasks/{task_id}/points/manual-batch", json={"text": "  \n"})).status_code == 400
    app.state.llm = StubLLM([json.dumps({
        "matches": [{"manual": manual[0]["tp_id"], "ai": ai_ids[0], "verdict": "一致", "reason": "均验证正常登录"},
                    {"manual": "TP999", "ai": ai_ids[0], "verdict": "一致"}],   # 非法 id / 重复占用被丢弃
        "manual_only": [{"tp_id": manual[1]["tp_id"], "reason": "AI 未覆盖锁定后找回密码"}],
        "ai_only": [{"tp_id": ai_ids[1], "reason": "属合理补充", "valuable": True}],
    }, ensure_ascii=False)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/compare")
    assert resp.status_code == 200, resp.text
    r = resp.json()["point_compare"]
    assert [m["manual"] for m in r["matched"]] == [manual[0]["tp_id"]] and r["matched"][0]["verdict"] == "一致"
    assert [m["tp_id"] for m in r["manual_only"]] == [manual[1]["tp_id"]] and "找回密码" in r["manual_only"][0]["reason"]
    assert {m["tp_id"] for m in r["ai_only"]} == set(ai_ids[1:]) and next(m for m in r["ai_only"] if m["tp_id"] == ai_ids[1])["valuable"] is True
    st = r["stats"]
    assert st["manual_total"] == 2 and st["ai_total"] == len(ai_ids) and st["matched"] == 1 and st["ai_recall"] == 0.5
    assert not r["degraded"]
    # 报告随任务持久化；再录入人工点后旧报告作废
    assert (await client.get(f"/api/v1/tasks/{task_id}")).json()["point_compare"]["stats"]["matched"] == 1
    await client.post(f"/api/v1/tasks/{task_id}/points/manual-batch", json={"text": "登录|记住密码"})
    assert (await client.get(f"/api/v1/tasks/{task_id}")).json()["point_compare"] is None
    # LLM 故障降级：文字相似配对 + degraded 标记
    class _Boom:
        async def chat(self, *a, **k):
            raise RuntimeError("down")
    app.state.llm = _Boom()
    r = (await client.post(f"/api/v1/tasks/{task_id}/points/compare")).json()["point_compare"]
    assert r["degraded"] and r["matched"] and r["matched"][0]["verdict"] == "待复核"


async def test_分批审核分批生成_完成后为新通过的测试点增量生成(client):
    data = await _create_confirm_task(client)
    task_id = data["task_id"]
    pts = [p for m in data["test_points"] for p in m["points"]]
    assert len(pts) >= 2
    first, second = pts[0], pts[1]
    # 第一批：只通过一条 → 生成
    await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [{"tp_id": first["tp_id"], "action": "approve"}]})
    app.state.llm = StubLLM([generator_reply(make_case(case_id="TC-登录-001", point_ids=[first["tp_id"]])), review_reply(True)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})
    assert resp.status_code == 200, resp.text
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert t["status"] == "completed" and len(t["result"]["cases"]) == 1
    assert t["context"]["generated_tp_ids"] == [first["tp_id"]] and t["points_pending_generation"] == []
    uid1 = t["result"]["cases"][0]["uid"]
    # 没有新通过的测试点：再点生成 409
    assert (await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})).status_code == 409
    # 第二批：完成后继续审核通过另一条 → 详情标出待生成 → 增量生成并入现有用例
    r = await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [{"tp_id": second["tp_id"], "action": "approve"}]})
    assert r.status_code == 200, r.text
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert t["points_pending_generation"] == [second["tp_id"]]
    # 新一批生成时模型仍从 001 编号；并入后按模块连续重编号，uid 不变
    app.state.llm = StubLLM([generator_reply(make_case(case_id="TC-登录-001", title="第二批用例", point_ids=[second["tp_id"]])), review_reply(True)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})
    assert resp.status_code == 200, resp.text
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert t["status"] == "completed" and len(t["result"]["cases"]) == 2
    assert t["result"]["cases"][0]["uid"] == uid1 and t["result"]["cases"][1]["title"] == "第二批用例"
    assert [c["case_id"] for c in t["result"]["cases"]] == ["TC-登录-001", "TC-登录-002"]
    assert set(t["context"]["generated_tp_ids"]) == {first["tp_id"], second["tp_id"]} and t["points_pending_generation"] == []
    assert any(x.get("action") == "增量生成" for x in t["result"]["trace"])


async def _wait_task(client, task_id):
    import asyncio
    for _ in range(100):
        await asyncio.sleep(0.03)
        t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
        if t["status"] in ("completed", "failed", "awaiting_confirmation"):
            return t
    raise AssertionError(f"任务未结束: {t['status']}")


async def test_确认生成可后台执行并轮询(client):
    data = await _create_confirm_task(client)
    task_id = data["task_id"]
    app.state.llm = StubLLM([generator_reply(make_case()), review_reply(True)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={"async_mode": True})
    assert resp.status_code == 200 and resp.json()["status"] == "queued" and resp.json()["poll_url"]
    t = await _wait_task(client, task_id)
    assert t["status"] == "completed" and len(t["result"]["cases"]) == 1
    assert t["analysis"]["test_points"] and t["context"]["generated_tp_ids"]  # 拆解留痕与分批生成记录都在
    assert "resume_stage" not in t["context"]
    # 后台失败（如模型返回异常）：状态恢复待确认并把原因留在任务上，可再次确认
    data2 = await _create_confirm_task(client)
    app.state.llm = StubLLM(["不是 JSON", "还不是 JSON"])
    assert (await client.post(f"/api/v1/tasks/{data2['task_id']}/confirm", json={"async_mode": True})).json()["status"] == "queued"
    t2 = await _wait_task(client, data2["task_id"])
    assert t2["status"] == "awaiting_confirmation" and t2["error"]


async def test_提案新增测试点不复用已删除编号_删空模块自动移除(client):
    data = await _create_confirm_task(client)  # TP001 TP002 TP003（TP003 为查漏新增）
    task_id = data["task_id"]
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [
        {"tp_id": "TP003", "action": "delete"},
        {"tp_id": "TP002", "action": "reject", "comment": "缺少锁定场景", "reject_types": ["场景遗漏"]},
    ]})
    assert resp.status_code == 200
    app.state.llm = StubLLM([json.dumps({"fixes": [], "additions": [{"module": "登录", "point": "连续错误5次锁定"}]},
                                        ensure_ascii=False)])
    assert (await client.post(f"/api/v1/tasks/{task_id}/points/fix")).status_code == 200
    resp = await client.post(f"/api/v1/tasks/{task_id}/fix/confirm", json={"accept_all": True})
    assert resp.status_code == 200
    ids = [p["tp_id"] for m in resp.json()["test_points"] for p in m["points"]]
    assert "TP003" not in ids and "TP004" in ids  # 删除的编号不复用
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert t["analysis"]["tp_seq"] == 4
    # 删掉模块里最后一个点：空模块不再保留
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/add", json={"points": [{"module": "独立模块", "point": "独立点"}]})
    tp = resp.json()["added"][0]["tp_id"]
    assert any(m["module"] == "独立模块" for m in resp.json()["test_points"])
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [{"tp_id": tp, "action": "delete"}]})
    assert not any(m["module"] == "独立模块" for m in resp.json()["test_points"])


async def test_手工新增相似测试点不再被静默跳过(client):
    data = await _create_confirm_task(client)
    task_id = data["task_id"]
    same = data["test_points"][0]["points"][0]["point"]
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/add", json={"points": [{"module": "登录", "point": same}]})
    assert resp.status_code == 200 and len(resp.json()["added"]) == 1 and resp.json()["skipped"] == []


# ---- 需求盲区闭环：转测试点 / 需求已说明 / 忽略 / 重开，质量摘要只数待处置 ----

ANALYST_REPLY_BLIND = json.dumps({
    "modules": [{"module": "登录", "points": [{"point": "正确账号密码登录成功", "dimension": "正常流程"}]}],
    "blind_spots": ["未说明锁定时长", "未说明找回密码入口", "未说明验证码有效期"],
}, ensure_ascii=False)


async def test_盲区闭环_处置状态与去向(client):
    app.state.llm = StubLLM([ANALYST_REPLY_BLIND, GAP_REPLY])
    resp = await client.post("/api/v1/tasks", data={"text": "登录需求", "confirm_points": "true"})
    assert resp.status_code == 200, resp.text
    task_id = resp.json()["task_id"]
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    spots = t["blind_spots"]
    assert [b["text"] for b in spots] == ["未说明锁定时长", "未说明找回密码入口", "未说明验证码有效期"]
    assert all(b["status"] == "open" and b["bs_id"].startswith("bs") for b in spots)
    a, b, c = (x["bs_id"] for x in spots)

    # 转测试点：以待审核测试点追加，来源 blind_spot，去向记 tp_id
    resp = await client.post(f"/api/v1/tasks/{task_id}/blind-spots/{a}",
                             json={"action": "to_point", "module": "登录", "dimension": "边界值"})
    assert resp.status_code == 200, resp.text
    added = resp.json()["added"][0]
    assert added["source"] == "blind_spot" and added["module"] == "登录" and added["status"] == "pending"
    assert added["point"] == "未说明锁定时长" and added["dimension"] == "边界值"
    assert resp.json()["blind_spot"]["status"] == "point" and resp.json()["blind_spot"]["target"]["tp_id"] == added["tp_id"]
    # 需求已说明：必须写明出处；已处置的不能再处置，重开后可以
    assert (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/{b}", json={"action": "explain"})).status_code == 400
    resp = await client.post(f"/api/v1/tasks/{task_id}/blind-spots/{b}", json={"action": "explain", "note": "第 3 节写了 30 分钟"})
    assert resp.json()["blind_spot"]["status"] == "explained" and resp.json()["blind_spot"]["note"] == "第 3 节写了 30 分钟"
    assert (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/{b}", json={"action": "ignore"})).status_code == 409
    resp = await client.post(f"/api/v1/tasks/{task_id}/blind-spots/{b}", json={"action": "reopen"})
    assert resp.json()["blind_spot"]["status"] == "open" and resp.json()["blind_spot"]["history"][0]["status"] == "explained"
    # 未关联需求的任务不能转待确认事项；非法动作 / 不存在的盲区
    resp = await client.post(f"/api/v1/tasks/{task_id}/blind-spots/{c}", json={"action": "to_question"})
    assert resp.status_code == 409 and "未关联需求" in resp.json()["detail"]
    assert (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/{c}", json={"action": "nope"})).status_code == 400
    assert (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/bs00000000", json={"action": "ignore"})).status_code == 404
    assert (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/{c}", json={"action": "ignore"})).status_code == 200

    # 列表接口：待处置在前
    spots = (await client.get(f"/api/v1/tasks/{task_id}/blind-spots")).json()["blind_spots"]
    assert [(x["bs_id"], x["status"]) for x in spots] == [(b, "open"), (a, "point"), (c, "ignored")]
    # 确认生成后质量摘要「待确认问题」只数待处置的 1 条，处置记录保留
    app.state.llm = StubLLM([generator_reply(make_case()), review_reply(True)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})
    assert resp.status_code == 200, resp.text
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert t["quality"]["待确认问题"] == 1
    assert {x["bs_id"]: x["status"] for x in t["blind_spots"]} == {a: "point", b: "open", c: "ignored"}
    await client.post(f"/api/v1/tasks/{task_id}/blind-spots/{b}", json={"action": "ignore"})
    assert (await client.get(f"/api/v1/tasks/{task_id}")).json()["quality"]["待确认问题"] == 0


def test_盲区实体_AI复核剔除项默认为需求已说明_人工可覆盖():
    from app.tasks.blind_spots import blind_spot_entities, open_blind_spots, record_blind_spot_action
    from app.tasks.store import TaskRecord

    rec = TaskRecord(task_id="t1", analysis={
        "test_points": [], "blind_spots": ["未说明找回密码入口"],
        "resolved_blind_spots": [{"text": "未说明锁定时长", "evidence": "第 3 片：30 分钟", "chunk": 3}],
    })
    ents = blind_spot_entities(rec)
    assert [(e["text"], e["status"], e.get("ai_resolved", False)) for e in ents] == [
        ("未说明找回密码入口", "open", False), ("未说明锁定时长", "explained", True)]
    assert "第 3 片已说明：第 3 片：30 分钟" in ents[1]["note"]
    assert open_blind_spots(rec) == ["未说明找回密码入口"]
    # 人工认为 AI 复核错了：重开 → 待处置
    record_blind_spot_action(rec, ents[1], "reopen", "tester")
    assert open_blind_spots(rec) == ["未说明找回密码入口", "未说明锁定时长"]


async def test_疑似重复对_忽略重开留痕_删除后自动已处置(client):
    analyst = json.dumps({"modules": [{"module": "登录", "points": [
        {"point": "正确账号密码登录成功", "dimension": "正常流程"},
        {"point": "正确账号密码登录成功后进入首页", "dimension": "正常流程"},
        {"point": "密码错误提示", "dimension": "异常流程"},
    ]}], "blind_spots": []}, ensure_ascii=False)
    dup = json.dumps({"duplicates": [{"a": "TP001", "b": "TP002", "verdict": "重复", "reason": "同一目标", "suggestion": "保留一条"}]}, ensure_ascii=False)
    app.state.llm = StubLLM([analyst, GAP_REPLY, dup])
    resp = await client.post("/api/v1/tasks", data={"text": "登录需求", "confirm_points": "true"})
    assert resp.status_code == 200, resp.text
    task_id = resp.json()["task_id"]
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    pairs = t["dup_pairs"]
    assert len(pairs) >= 1 and pairs[0]["status"] == "open" and pairs[0]["key"] == "TP001|TP002"
    key = pairs[0]["key"]
    # 忽略 → 已忽略（带处置人）；重复忽略 409；重开 → 待处置
    r = (await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/{key}", json={"action": "ignore", "note": "都要保留"})).json()
    assert r["dup_pair"]["status"] == "ignored" and r["dup_pair"]["note"] == "都要保留" and r["dup_pair"]["by"] == "anonymous"
    assert (await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/{key}", json={"action": "ignore"})).status_code == 409
    assert (await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/{key}", json={"action": "nope"})).status_code == 400
    assert (await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/none", json={"action": "ignore"})).status_code == 404
    r = (await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/{key}", json={"action": "reopen"})).json()
    assert r["dup_pair"]["status"] == "open"
    # 重新检查后忽略记录保留；删除其中一条后自动已处置且不可重开
    await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/{key}", json={"action": "ignore"})
    app.state.llm = StubLLM([dup])
    r = (await client.post(f"/api/v1/tasks/{task_id}/points/dup-check")).json()
    assert next(p for p in r["dup_pairs"] if p["key"] == key)["status"] == "ignored"
    await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [{"tp_id": "TP002", "action": "delete"}]})
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    p = next(p for p in t["dup_pairs"] if p["key"] == key)
    assert p["status"] == "resolved" and p["deleted"] == ["TP002"]
    assert (await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/{key}", json={"action": "reopen"})).status_code == 409


async def test_盲区与重复对_批量处置(client):
    app.state.llm = StubLLM([ANALYST_REPLY_BLIND, GAP_REPLY])
    resp = await client.post("/api/v1/tasks", data={"text": "登录需求", "confirm_points": "true"})
    task_id = resp.json()["task_id"]
    spots = (await client.get(f"/api/v1/tasks/{task_id}")).json()["blind_spots"]
    a, b, c = (x["bs_id"] for x in spots)
    # 批量说明必须带出处；批量忽略：已处置的跳过并带原因，其余生效，一次落盘
    assert (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/batch", json={"action": "explain", "ids": [a]})).status_code == 200
    r = (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/batch", json={"action": "explain", "ids": [a], "note": "3.2 节"})).json()
    assert r["skipped"] == [] and r["applied"][0]["status"] == "explained"
    r = (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/batch", json={"action": "ignore", "ids": [a, b, c, "bs404"]})).json()
    assert [x["status"] for x in r["applied"]] == ["ignored", "ignored"]
    assert {x["id"]: x["reason"] for x in r["skipped"]} == {a: "该盲区已处置（explained），请先重开", "bs404": "盲区不存在"}
    assert {x["bs_id"]: x["status"] for x in r["blind_spots"]} == {a: "explained", b: "ignored", c: "ignored"}
    assert (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/batch", json={"action": "ignore", "ids": []})).status_code == 400
    assert (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/batch", json={"action": "xx", "ids": [a]})).status_code == 400
    # 批量重开后批量转测试点：每条用自己的原文，模块统一
    r = (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/batch", json={"action": "reopen", "ids": [b, c]})).json()
    assert [x["status"] for x in r["applied"]] == ["open", "open"]
    r = (await client.post(f"/api/v1/tasks/{task_id}/blind-spots/batch",
                           json={"action": "to_point", "ids": [b, c], "module": "登录", "point": "不应生效的统一描述"})).json()
    assert [x["status"] for x in r["applied"]] == ["point", "point"]
    pts = (await client.get(f"/api/v1/tasks/{task_id}")).json()["analysis"]["test_points"]
    new_pts = [p for m in pts for p in m["points"] if p.get("source") == "blind_spot"]
    assert sorted(p["point"] for p in new_pts) == sorted([spots[1]["text"], spots[2]["text"]])  # point 字段不生效，各用原文
    assert all(m["module"] == "登录" for m in pts if any(p.get("source") == "blind_spot" for p in m["points"]))
    # 重复对批量忽略 / 重开
    app.state.llm = StubLLM([json.dumps({"pairs": []})])
    pairs = (await client.post(f"/api/v1/tasks/{task_id}/points/dup-check")).json().get("dup_pairs") or []
    if pairs:
        keys = [p["key"] for p in pairs]
        r = (await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/batch", json={"action": "ignore", "keys": keys + ["nokey"]})).json()
        assert len(r["applied"]) == len(keys) and r["skipped"] == [{"key": "nokey", "reason": "疑似重复对不存在"}]
        r = (await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/batch", json={"action": "reopen", "keys": keys})).json()
        assert all(p["status"] == "open" for p in r["applied"])
    assert (await client.post(f"/api/v1/tasks/{task_id}/points/dup-pairs/batch", json={"action": "ignore", "keys": []})).status_code == 400


async def test_删除非空项目_默认拒绝并列出内容_cascade连同需求计划任务一起删(client):
    """项目不空时默认 400 并说清里面有什么；contents 接口给清单；cascade=true 连同内容一起删：需求彻底清除（回收站里的也清）、
    计划连执行轮次一起删、任务进回收站；之后项目 404、任务列表与需求列表都空。"""
    from tests.stubs import ANALYST_REPLY

    await client.post("/api/v1/projects", json={"name": "要整删"})
    app.state.llm = StubLLM([ANALYST_REPLY, generator_reply(make_case()), review_reply(True)])
    task_id = (await client.post("/api/v1/tasks", data={"text": "登录需求", "project": "要整删"})).json()["task_id"]
    r1 = (await client.post("/api/v1/requirements", data={"project": "要整删", "title": "需求一", "text": "原文一"})).json()["req_id"]
    r2 = (await client.post("/api/v1/requirements", data={"project": "要整删", "title": "需求二", "text": "原文二"})).json()["req_id"]
    assert (await client.delete(f"/api/v1/requirements/{r2}")).status_code == 200   # 一条进回收站
    plan_id = (await client.post("/api/v1/plans", json={"name": "计划A", "project": "要整删"})).json()["plan_id"]

    c = (await client.get("/api/v1/projects/要整删/contents")).json()
    assert c["tasks"] == 1 and c["plans"] == 1 and c["requirements"] == 1 and c["recycled_requirements"] == 1 and c["running"] == 0 and not c["empty"]
    resp = await client.delete("/api/v1/projects/要整删")
    assert resp.status_code == 400 and "1 个任务、1 个测试计划、1 条需求" in resp.json()["detail"] and "cascade=true" in resp.json()["detail"]
    assert (await client.get("/api/v1/projects/要整删/versions")).status_code == 200   # 没删

    resp = await client.delete("/api/v1/projects/要整删", params={"cascade": "true"})
    assert resp.status_code == 200 and resp.json() == {"deleted": "要整删", "cascade": True, "tasks": 1, "plans": 1, "requirements": 1}
    assert (await client.get("/api/v1/projects/要整删/contents")).status_code == 404
    assert (await client.get(f"/api/v1/requirements/{r1}")).status_code == 404 and (await client.get(f"/api/v1/requirements/{r2}")).status_code == 404
    assert (await client.get(f"/api/v1/plans/{plan_id}")).status_code == 404
    assert (await client.get("/api/v1/tasks", params={"project": "要整删"})).json()["tasks"] == []
    assert any(t["task_id"] == task_id for t in (await client.get("/api/v1/tasks", params={"include_deleted": "true"})).json()["tasks"])   # 任务在回收站
    assert (await client.get("/api/v1/projects/不存在/contents")).status_code == 404


# ---------- 查重候选按页面分档 / 同模块同文去重 / 裁判分批带模块 ----------


def test_查重候选_同页面按低阈值_不同页面文字几乎一样才算_候选带模块():
    from app.tasks.points import module_page

    mid = "选择真实会员时仅返回真实会员和机器人记录"   # 与「选择真实会员时仅返回真实会员记录」相似度在 0.45–0.7 之间
    modules = assign_entities([
        {"module": "后台-日榜排名-字段说明", "points": ["商户名称取会员所属商户名称", "选择真实会员时仅返回真实会员记录"]},
        {"module": "后台-日榜排名-查询功能说明", "points": [mid]},
        {"module": "后台-月榜排名-字段说明", "points": ["商户名称取会员所属商户名称", mid]},
    ])
    found = {(p["a"], p["b"]): p for p in duplicate_candidates(modules)}
    # 同页面（日榜排名）不同区块：中等相似按 0.45 进候选，并带两边模块
    assert found[("TP002", "TP003")]["a_module"] == "后台-日榜排名-字段说明" and found[("TP002", "TP003")]["b_module"] == "后台-日榜排名-查询功能说明"
    # 不同页面：同文的「商户名称…」进（1.0 ≥ 0.7）；同样中等相似的那对不进
    assert found[("TP001", "TP004")]["b_module"] == "后台-月榜排名-字段说明"
    assert ("TP002", "TP005") not in found
    # 月榜页内两条同文（TP003 与 TP005 不同页面但同文）→ 进
    assert ("TP003", "TP005") in found
    assert module_page("后台-日榜排名-字段说明") == "后台-日榜排名" and module_page("活动运营后台-日榜统计") == "活动运营后台" and module_page("登录") == "登录"


def test_模块所属页面按父子关系定_需求前缀下的页面级模块不并成同一页():
    from app.tasks.points import module_pages

    names = ["后台-日榜排名", "后台-日榜排名-字段说明", "后台-月榜排名", "后台-月榜排名-总计",
             "后台-API接口-会员流水", "前端-PC-活动入口说明", "前端-PC-活动入口说明-报名规则说明", "登录-密码", "登录-验证码"]
    pages = module_pages(names)
    # 有子模块的自己就是页面：「后台-日榜排名」与「后台-月榜排名」是两个页面，不会都算成「后台」
    assert pages["后台-日榜排名"] == "后台-日榜排名" and pages["后台-月榜排名"] == "后台-月榜排名"
    assert pages["后台-日榜排名-字段说明"] == "后台-日榜排名" and pages["后台-月榜排名-总计"] == "后台-月榜排名"
    assert pages["前端-PC-活动入口说明"] == "前端-PC-活动入口说明" == pages["前端-PC-活动入口说明-报名规则说明"]
    # 没有父模块也没有子模块：按名字去掉最后一段
    assert pages["后台-API接口-会员流水"] == "后台-API接口" and pages["登录-密码"] == pages["登录-验证码"] == "登录"
    # 查重候选随之：日榜页与月榜页的页面级同文测试点按跨页面阈值（1.0 进）、中等相似不进
    modules = assign_entities([
        {"module": "后台-日榜排名", "points": ["总人数为当前查询结果按会员账号去重后的数量", "切换商户币种后总计金额按对应币种更新"]},
        {"module": "后台-日榜排名-字段说明", "points": ["会员昵称取会员当前昵称"]},
        {"module": "后台-月榜排名", "points": ["总人数为当前查询结果按会员账号去重后的数量", "切换商户币种后列表奖励金额按对应币种显示"]},
        {"module": "后台-月榜排名-字段说明", "points": ["会员昵称取会员当前昵称"]},
    ])
    found = {(p["a"], p["b"]) for p in duplicate_candidates(modules)}
    assert ("TP001", "TP004") in found and ("TP003", "TP006") in found
    assert ("TP002", "TP005") not in found


def test_拆解结果同模块同文去重_同名模块并一个_不同模块同文不动():
    from app.tasks.points import dedupe_module_points

    modules = [
        {"module": "日榜排名-字段说明", "points": [{"point": "会员昵称取会员当前昵称。", "dimension": "数据"}, {"point": "连赢次数取最高连赢次数", "dimension": "数据"}]},
        {"module": "日榜排名-字段说明", "points": [{"point": "会员昵称取会员当前昵称", "dimension": "数据"}, {"point": "积分取当日积分", "dimension": "数据"}]},
        {"module": "月榜排名-字段说明", "points": [{"point": "会员昵称取会员当前昵称。", "dimension": "数据"}]},
    ]
    assert dedupe_module_points(modules) == 1
    assert [m["module"] for m in modules] == ["日榜排名-字段说明", "月榜排名-字段说明"]
    assert [p["point"] for p in modules[0]["points"]] == ["会员昵称取会员当前昵称。", "连赢次数取最高连赢次数", "积分取当日积分"]
    assert len(modules[1]["points"]) == 1


async def test_重复复核分批_每批带模块_没判到的待复核_一批故障只影响该批():
    from app.agents import quality
    from app.agents.quality import run_dup_judge

    pairs = [{"a": f"TP{i:03d}", "a_point": f"规则 {i}", "a_module": "日榜排名-字段说明",
              "b": f"TP{i + 100:03d}", "b_point": f"规则 {i} 的另一种说法", "b_module": "月榜排名-字段说明",
              "similarity": 0.8, "reason": "文字高度相似"} for i in range(1, 6)]
    old = quality.DUP_JUDGE_BATCH
    quality.DUP_JUDGE_BATCH = 2
    try:
        r1 = json.dumps({"duplicates": [{"a": "TP001", "b": "TP101", "verdict": "不重复", "reason": "日榜与月榜两个对象"},
                                        {"a": "TP002", "b": "TP102", "verdict": "重复", "reason": "同一规则"}]}, ensure_ascii=False)
        r2 = "这不是 JSON"   # 第二批解析失败（_chat_json 重试一次仍失败）
        r3 = json.dumps({"duplicates": [{"a": "TP005", "b": "TP105", "verdict": "不重复", "reason": "对象不同"}]}, ensure_ascii=False)
        llm = StubLLM([r1, r2, r2, r3])
        out = await run_dup_judge(llm, "需求原文", pairs, None)
    finally:
        quality.DUP_JUDGE_BATCH = old
    assert len(llm.calls) == 4   # 3 批 + 第二批重试一次
    user = llm.calls[0]["messages"][-1]["content"]
    assert "日榜排名-字段说明" in user and "月榜排名-字段说明" in user and "TP003" not in user   # 每批只带本批的对、带模块
    by = {(p["a"], p["b"]): p for p in out}
    assert ("TP001", "TP101") not in by and ("TP005", "TP105") not in by          # 不重复的不再打扰人工
    assert by[("TP002", "TP102")]["verdict"] == "重复"
    assert by[("TP003", "TP103")]["verdict"] == "待复核" and "复核失败" in by[("TP003", "TP103")]["suggestion"]  # 故障批
    assert by[("TP004", "TP104")]["verdict"] == "待复核"


async def test_重复复核_同文却判不重复的留在列表标同文不同对象_文字不同的不重复才去掉():
    from app.agents.quality import run_dup_judge

    pairs = [
        {"a": "TP001", "a_point": "商户名称取会员所属商户名称。", "a_module": "后台-日榜排名-字段说明",
         "b": "TP201", "b_point": "商户名称取会员所属商户名称", "b_module": "后台-月榜排名-字段说明", "similarity": 1.0, "reason": "文字高度相似"},
        {"a": "TP002", "a_point": "跨度超过 90 天不允许查询", "a_module": "后台-日榜排名-查询功能说明",
         "b": "TP202", "b_point": "跨度超过 3 个月不允许查询", "b_module": "后台-月榜排名-查询功能说明", "similarity": 0.75, "reason": "文字高度相似"},
    ]
    llm = StubLLM([json.dumps({"duplicates": [
        {"a": "TP001", "b": "TP201", "verdict": "不重复", "reason": "日榜与月榜两个页面各验一次"},
        {"a": "TP002", "b": "TP202", "verdict": "不重复", "reason": "日榜按天、月榜按月，条件不同"},
    ]}, ensure_ascii=False)])
    out = await run_dup_judge(llm, "需求原文", pairs, None)
    assert [p["a"] for p in out] == ["TP001"]
    assert out[0]["verdict"] == "同文不同对象" and "各验一次" in out[0]["reason"] and "忽略" in out[0]["suggestion"]
    system = llm.calls[0]["messages"][0]["content"]
    assert "合写复述算重复" in system


# ---------- 盲区出处：模型给的 module / ref 优先，没有的按章节号在整理稿里反查 ----------


def _整理稿_for_bs():
    return "\n".join([
        "【需求 1/2：后台】", "# 「后台」需求整理稿", "", "## 3. 商户管理后台", "",
        "### 3.4 日榜排名 〔来源：x.zip › 商户管理后台 / 日榜排名〕", "", "#### 3.4.5 字段说明", "- 会员昵称", "#### 3.4.7 查询功能说明", "- 按天排重",
        "【需求 2/2：前端】", "# 「前端」需求整理稿", "", "## 3. PC", "", "### 3.4 PC-活动规则 〔来源：y.zip › PC / PC-活动规则〕", "", "#### 3.4.7 规则弹窗", "- 关闭",
    ])


def test_盲区出处_模型给的优先_没有的按章节号反查_多需求按模块前缀区分():
    from app.agents.graph import _blind_spot_refs, _blind_spots
    from app.tasks.blind_spots import blind_spot_entities, locate_blind_spot
    from app.tasks.store import TaskRecord

    data = {"modules": [], "blind_spots": [
        {"text": "排重键未说明", "module": "后台-日榜排名-查询功能说明", "ref": "后台 › 3.4.7 查询功能说明 › 按天排重"},
        "3.4.5 会员昵称取值时点未说明",
        {"description": "只有文本的对象形态"},
    ]}
    assert _blind_spots(data) == ["排重键未说明", "3.4.5 会员昵称取值时点未说明", "只有文本的对象形态"]
    assert _blind_spot_refs(data) == {"排重键未说明": {"module": "后台-日榜排名-查询功能说明", "ref": "后台 › 3.4.7 查询功能说明 › 按天排重"}}

    req = _整理稿_for_bs()
    modules = [{"module": "后台-日榜排名-字段说明", "points": []}, {"module": "后台-日榜排名-查询功能说明", "points": []},
               {"module": "前端-PC-活动规则", "points": []}]
    # 章节号 3.4.5 → 祖先链 + 名字含页面 / 区块标题的模块
    assert locate_blind_spot("3.4.5 会员昵称取值时点未说明", req, modules) == {
        "module": "后台-日榜排名-字段说明", "ref": "后台 › 3 商户管理后台 › 3.4 日榜排名 › 3.4.5 字段说明"}
    # 两份整理稿都有 3.4.7：默认取第一份；给了模块前缀「前端-」则取前端那份
    assert locate_blind_spot("3.4.7 排重键未说明", req, modules)["ref"].startswith("后台 › ")
    assert locate_blind_spot("3.4.7 弹窗关闭后状态未说明", req, modules, hint_module="前端-PC-活动规则")["ref"] == "前端 › 3 PC › 3.4 PC-活动规则 › 3.4.7 规则弹窗"
    assert locate_blind_spot("没有章节号的盲区", req, modules) == {}

    rec = TaskRecord(task_id="t1", context={"requirement": req}, analysis={
        "test_points": modules, "blind_spots": ["排重键未说明", "3.4.5 会员昵称取值时点未说明", "没有章节号的盲区"],
        "blind_spot_refs": _blind_spot_refs(data)})
    ents = {e["text"]: e for e in blind_spot_entities(rec)}
    assert ents["排重键未说明"]["module"] == "后台-日榜排名-查询功能说明" and ents["排重键未说明"]["ref"].endswith("按天排重")
    assert ents["3.4.5 会员昵称取值时点未说明"]["ref"] == "后台 › 3 商户管理后台 › 3.4 日榜排名 › 3.4.5 字段说明"
    assert "ref" not in ents["没有章节号的盲区"] and "module" not in ents["没有章节号的盲区"]


def _整理稿_for_gap(区块字数=900):
    big = "- 规则一条，带数值 100 与条件。\n" * (区块字数 // 18 + 1)
    return "\n".join([
        "# 「后台」需求整理稿", "", "## 1. 需求说明", "", "### 1.1 创建活动", big,
        "## 2. 活动运营后台", "", "### 2.1 活动列表", "", "#### 2.1.1 基本配置", big, "#### 2.1.2 数据配置", big,
        "### 2.2 派奖查询", "", "#### 2.2.1 派奖查询", big,
    ])


async def test_查漏按章节分批_各批只带本批章节的测试点_矩阵合并_一批故障只丢该批(monkeypatch):
    """线上 7.6 万字需求 + 1255 条测试点一次送查漏（10.6 万 token），模型 9 秒回「全已覆盖、新增 0 条」，整轮等于没跑。"""
    import app.agents.quality as Q
    from app.agents.outline import requirement_chunks

    monkeypatch.setattr(Q, "GAP_BATCH_CHARS", 1500)
    req = _整理稿_for_gap()
    n = len(requirement_chunks(req, 1500)[0])
    assert n >= 3
    modules = [
        {"module": "需求说明-创建活动", "points": [{"tp_id": "TP001", "point": "活动名称必填", "dimension": "输入校验", "ref": "1.1 创建活动 › 名称"}]},
        {"module": "活动列表-基本配置", "points": [{"tp_id": "TP002", "point": "奖池比例 75%", "dimension": "数据", "ref": "2.1.1 基本配置 › 比例"}]},
        {"module": "活动运营后台-派奖查询", "points": [{"tp_id": "TP003", "point": "待派可查", "dimension": "正常流程", "ref": "2.2.1 派奖查询 › 状态"}]},
        {"module": "模型自拟的名字", "points": [{"tp_id": "TP004", "point": "按 ref 的章节号归批", "dimension": "数据", "ref": "2.1.2 数据配置 › 一条"}]},
        {"module": "哪都对不上", "points": [{"tp_id": "TP005", "point": "每批都带", "dimension": "数据", "ref": "无"}]},
    ]
    batches = Q.gap_batches(req, modules)
    assert len(batches) == n
    where = {p["tp_id"]: [i for i, (_, ms) in enumerate(batches) if any(pp["tp_id"] == p["tp_id"] for m in ms for pp in m["points"])]
             for m in modules for p in m["points"]}
    chunk_of = lambda s: [i for i, (t, _) in enumerate(batches) if s in t]  # noqa: E731
    assert where["TP001"] == chunk_of("### 1.1 创建活动") and where["TP002"] == chunk_of("#### 2.1.1 基本配置")
    assert where["TP003"] == chunk_of("#### 2.2.1 派奖查询") and where["TP004"] == chunk_of("#### 2.1.2 数据配置")
    assert where["TP005"] == list(range(n))   # 模块名与 ref 都对不上：每批都带

    ok = lambda cov, adds: json.dumps({"coverage": cov, "additions": adds}, ensure_ascii=False)  # noqa: E731
    full = {d: "已覆盖" for d in Q.DIMENSIONS}
    replies = [ok({**full, "输入校验": "未覆盖", "网络": "不适用"}, [{"module": "需求说明-创建活动", "point": "名称超长不可提交", "dimension": "输入校验", "ref": "1.1 创建活动 › 名称"}]),
               "坏掉的 JSON", "还是坏的",   # 第二批两次都解析失败 → 该批丢弃
               ok({**full, "网络": "不适用"}, [{"module": "活动运营后台-派奖查询", "point": "驳回状态可查", "dimension": "正常流程", "ref": "2.2.1 派奖查询 › 状态"}])]
    replies += [ok({**full, "网络": "不适用"}, [])] * (n - 3)
    llm = StubLLM(replies)
    gap = await Q.run_gap_check(llm, req, modules, None)
    assert gap["batches"] == n and len(llm.calls) == n + 1
    assert gap["coverage"]["输入校验"] == "未覆盖" and gap["coverage"]["正常流程"] == "已覆盖" and gap["coverage"]["网络"] == "不适用"
    assert [a["point"] for a in gap["additions"]] == ["名称超长不可提交", "驳回状态可查"]
    first = llm.calls[0]["messages"][-1]["content"]
    assert "第 1/" in first and "活动名称必填" in first and "奖池比例 75%" not in first and "每批都带" in first
    # 全部对不上（需求不是整理稿、模块名模型自拟）→ 不分批，整篇一次
    plain = "登录需求：\n" + "账号密码登录，连续错误 5 次锁定。\n" * 300
    assert Q.gap_batches(plain, [{"module": "登录", "points": [{"tp_id": "TP001", "point": "正常登录"}]}]) == [(plain, [{"module": "登录", "points": [{"tp_id": "TP001", "point": "正常登录"}]}])]
    assert Q.merge_coverage([{"网络": "不适用"}, {"网络": "不适用"}])["网络"] == "不适用"
    assert Q.merge_coverage([{"网络": "不适用"}, {"网络": "待确认"}])["网络"] == "待确认"


# ---- 经验延伸：测试点页手动再跑一轮 / 拆解时自动跑 ----


async def test_points_extend_手动经验延伸_只新增标来源与依据_确认后依据进生成_导出带依据列(client, tmp_path):
    from openpyxl import load_workbook

    from tests.test_agents import EXTEND_REPLY

    data = await _create_confirm_task(client)
    task_id = data["task_id"]
    app.state.llm = StubLLM([EXTEND_REPLY])
    resp = await client.post(f"/api/v1/tasks/{task_id}/points/extend")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    # 明文 2 条 + 查漏 1 条 = 3 条已有 → 上限 max(3, 1) = 3；复述 / 没依据 / 模块对不上的跳过
    assert [p["source"] for p in body["added"]] == ["derived"] * 3
    assert all(p["basis"] and p["status"] == "pending" and p["tp_id"].startswith("TP") for p in body["added"])
    reasons = "｜".join(s["reason"] for s in body["skipped"])
    assert "没有写为什么要测" in reasons and "模块不在已有测试点里" in reasons and "超出延伸上限" in reasons
    added_ids = [p["tp_id"] for p in body["added"]]
    # 任务详情里带 basis，前端据此显示「延伸」与依据
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    pts = {p["tp_id"]: p for m in t["analysis"]["test_points"] for p in m["points"]}
    assert pts[added_ids[0]]["source"] == "derived" and pts[added_ids[0]]["basis"].startswith("弱网")
    # 导出：来源列「经验延伸」，末列延伸依据
    r = await client.get(f"/api/v1/tasks/{task_id}/points/export")
    f = tmp_path / "pts.xlsx"
    f.write_bytes(r.content)
    rows = [[str(c) if c is not None else "" for c in row] for row in load_workbook(str(f)).active.iter_rows(values_only=True)]
    assert rows[0][-1] == "延伸依据"
    by_id = {row[0]: row for row in rows[1:]}
    assert by_id[added_ids[0]][5] == "经验延伸" and by_id[added_ids[0]][-1].startswith("弱网")
    assert by_id["TP001"][-1] == ""
    # 通过一条明文 + 一条延伸后确认生成：延伸点连依据一起进生成 Prompt
    await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [
        {"tp_id": "TP001", "action": "approve"}, {"tp_id": added_ids[0], "action": "approve"}]})
    app.state.llm = StubLLM([generator_reply(make_case()), review_reply(True)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})
    assert resp.status_code == 200, resp.text
    gen_prompt = app.state.llm.calls[0]["messages"][1]["content"]
    assert "登录请求发出后断网再恢复" in gen_prompt and "弱网下请求重发导致重复登录是常见缺陷" in gen_prompt


async def test_拆解时自动经验延伸_开关打开时拆完明文紧接着延伸再查漏(client, monkeypatch):
    from app.config import get_settings
    from tests.test_agents import EXTEND_REPLY

    monkeypatch.setattr(get_settings(), "point_extension", True)
    app.state.llm = StubLLM([ANALYST_REPLY_V2, EXTEND_REPLY, GAP_REPLY])
    resp = await client.post("/api/v1/tasks", data={"text": "登录需求：账号密码登录，连续错误5次锁定", "confirm_points": "true"})
    assert resp.status_code == 200, resp.text
    points = resp.json()["test_points"][0]["points"]
    srcs = [p["source"] for p in points]
    assert srcs == ["ai", "ai", "derived", "derived", "derived", "gap"]   # 明文 → 延伸（上限 3）→ 查漏新增
    assert all(p["basis"] for p in points if p["source"] == "derived") and len({p["tp_id"] for p in points}) == 6
    ext_user = app.state.llm.calls[1]["messages"][-1]["content"]
    assert "正确账号密码登录成功" in ext_user   # 延伸看到了全部明文点
