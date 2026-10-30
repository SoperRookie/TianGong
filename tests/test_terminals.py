"""按端（H5 / 游戏后台管理 / 商户后台）处理：整理稿三端汇总、测试点归端与按端单独生成、端到端场景、按端导出。"""

import json
from urllib.parse import unquote

import httpx
import pytest
from asgi_lifespan import LifespanManager
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from tests.stubs import StubLLM, generator_reply, make_case, review_reply


@pytest.fixture
async def client():
    async with LifespanManager(app):
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
            yield c


ANALYST_3_TERMINALS = json.dumps({
    "modules": [
        {"module": "活动运营后台-活动列表", "points": [{"point": "活动名称必填，空值不可保存", "dimension": "输入校验", "ref": "活动列表 › 基本配置"},
                                              {"point": "活动结束后不可再编辑奖励配置", "dimension": "状态转换", "ref": "活动列表 › 编辑"}]},
        {"module": "商户管理后台-日榜排名", "points": [{"point": "日榜按连赢次数降序排列", "dimension": "数据", "ref": "日榜排名 › 排序"}]},
        {"module": "H5-活动入口", "points": [{"point": "未报名会员点击入口弹出报名弹窗", "dimension": "正常流程", "ref": "活动入口 › 报名"}]},
    ],
    "blind_spots": [],
}, ensure_ascii=False)
NO_GAP = json.dumps({"coverage": {}, "additions": []}, ensure_ascii=False)
E2E_REPLY = json.dumps({"scenarios": [
    {"name": "后台配置 → 会员报名 → 商户报表", "terminals": ["game_admin", "h5", "merchant_admin"],
     "flow": "游戏后台：配置活动 → H5：会员报名 → 商户后台：报名记录入账",
     "points": [{"point": "后台保存的报名费与 H5 报名弹窗显示金额一致", "dimension": "数据", "ref": "活动列表 › 报名费", "basis": "跨端金额口径"},
                {"point": "H5 报名成功后商户后台报名记录立即出现该会员", "dimension": "数据", "ref": "报名记录 › 列表", "basis": "跨端数据流转"}]},
    {"name": "单端的不算", "terminals": ["h5"], "flow": "H5：报名", "points": [{"point": "x", "dimension": "数据", "ref": "a", "basis": "b"}]},
]}, ensure_ascii=False)


def test_关键词判端_模块名按章节对回端_用例按来源测试点归端():
    from app.terminals import case_terminal, classify_name, label_of, terminal_of_module

    assert classify_name("活动运营后台") == "game_admin" and classify_name("商户管理后台") == "merchant_admin"
    assert classify_name("H5 / 活动入口") == "h5" and classify_name("需求说明") == "common" and classify_name("连赢活动统计") is None
    assert classify_name("游戏后台-大厅配置") == "game_admin"   # 后台类先判，不被「大厅」抢走
    secs = {"连赢活动统计": "game_admin", "商户管理后台": "merchant_admin"}
    assert terminal_of_module("连赢活动统计-日榜统计", secs) == "game_admin"
    assert terminal_of_module("商户管理后台-日榜排名-查询与排序", secs) == "merchant_admin"
    assert terminal_of_module("需求A-连赢活动统计-日榜统计", {"需求A-连赢活动统计": "game_admin"}, ["需求A"]) == "game_admin"
    assert terminal_of_module("没人认识的模块", secs) == "other"
    assert label_of("merchant_admin") == "商户后台" and label_of(None) == "其他"
    assert case_terminal({"point_ids": ["TP001", "TP002"], "module": "X"}, {"TP001": "h5", "TP002": "h5"}, {}) == "h5"
    assert case_terminal({"point_ids": [], "module": "商户管理后台-报名记录"}, {}, {}) == "merchant_admin"
    assert case_terminal({"point_ids": [], "module": "X"}, {}, {"X": "game_admin"}) == "game_admin"


def test_整理稿三端汇总_按顶层章节归端_人工覆盖优先():
    from app.terminals import section_terminals, terminal_summary

    outline = [
        {"no": "1", "gno": "", "group": "", "title": "需求说明", "blocks": [{"rows": [{"t": "item", "id": "a"}, {"t": "item", "id": "b"}]}]},
        {"no": "2.1", "gno": "2", "group": "活动运营后台", "title": "活动列表", "blocks": [{"rows": [{"t": "table", "ids": ["c", "d", "e"]}]}, {"rows": [{"t": "item", "id": "f"}]}]},
        {"no": "2.2", "gno": "2", "group": "活动运营后台", "title": "派奖查询", "blocks": [{"rows": [{"t": "item", "id": "g"}]}]},
        {"no": "3", "gno": "", "group": "", "title": "连赢活动统计", "blocks": [{"rows": [{"t": "item", "id": "h"}]}]},
        {"no": "4.1", "gno": "4", "group": "商户管理后台", "title": "日榜排名", "blocks": [{"rows": [{"t": "example", "id": "i", "ids": ["j", "k"]}]}]},
    ]
    terms = section_terminals(outline)
    assert terms == {"需求说明": "common", "活动运营后台": "game_admin", "连赢活动统计": "other", "商户管理后台": "merchant_admin"}
    terms = section_terminals(outline, {"连赢活动统计": "merchant_admin", "需求说明": "bogus"})   # 人工 / 模型给的优先；无效的忽略
    assert terms["连赢活动统计"] == "merchant_admin" and terms["需求说明"] == "common"
    s = terminal_summary(outline, {"连赢活动统计": "merchant_admin"})
    assert [t["key"] for t in s["terminals"]] == ["game_admin", "merchant_admin", "common"]
    game = s["terminals"][0]
    assert game["pages"] == 2 and game["blocks"] == 3 and game["items"] == 5 and [p["title"] for p in game["sections"][0]["pages"]] == ["活动列表", "派奖查询"]
    merchant = s["terminals"][1]
    assert [x["key"] for x in merchant["sections"]] == ["连赢活动统计", "商户管理后台"] and merchant["items"] == 4   # 示例图：说明 + 两条文案
    assert s["total"] == {"sections": 4, "pages": 5, "blocks": 6, "items": 11}


async def _three_terminal_task(client):
    app.state.llm = StubLLM([ANALYST_3_TERMINALS, NO_GAP])
    resp = await client.post("/api/v1/tasks", data={"text": "连赢挑战赛需求：活动运营后台、商户管理后台、H5", "confirm_points": "true"})
    assert resp.status_code == 200, resp.text
    return resp.json()["task_id"]


def _rows(content: bytes, tmp_path, name="f.xlsx"):
    f = tmp_path / name
    f.write_bytes(content)
    return [[str(c) if c is not None else "" for c in row] for row in load_workbook(str(f)).active.iter_rows(values_only=True)]


async def test_测试点按模块名归端_按端单独确认生成_其余端随后增量_按端导出(client, tmp_path):
    task_id = await _three_terminal_task(client)
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    pts = {p["tp_id"]: dict(p, module=m["module"]) for m in t["analysis"]["test_points"] for p in m["points"]}
    by_term = {}
    for p in pts.values():
        by_term.setdefault(p["terminal"], []).append(p["tp_id"])
    assert set(by_term) == {"game_admin", "merchant_admin", "h5"} and len(by_term["game_admin"]) == 2
    assert [x["key"] for x in t["point_terminals"]] == ["h5", "game_admin", "merchant_admin"]   # 按端清单顺序
    assert {x["key"]: x["total"] for x in t["point_terminals"]} == {"h5": 1, "game_admin": 2, "merchant_admin": 1}
    assert [x["key"] for x in t["terminals"]][:3] == ["h5", "game_admin", "merchant_admin"]
    # 商户后台与游戏后台的点都审完，H5 还没审
    items = [{"tp_id": tp, "action": "approve"} for tp in by_term["merchant_admin"] + by_term["game_admin"]]
    assert (await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": items})).status_code == 200
    # 未知的端 400；只为商户后台生成
    assert (await client.post(f"/api/v1/tasks/{task_id}/confirm", json={"terminal": "pc"})).status_code == 400
    m_tp = by_term["merchant_admin"][0]
    app.state.llm = StubLLM([generator_reply(make_case(case_id="TC-日榜-001", module="商户管理后台-日榜排名", title="日榜降序", point_ids=[m_tp])), review_reply(True)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={"terminal": "merchant_admin"})
    assert resp.status_code == 200, resp.text
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert t["status"] == "completed" and [c["title"] for c in t["result"]["cases"]] == ["日榜降序"]
    assert t["context"]["generated_tp_ids"] == [m_tp] and "confirm_terminal" not in t["context"]
    assert set(t["points_pending_generation"]) == set(by_term["game_admin"])   # 游戏后台已通过、待生成；H5 没审不算
    stats = {x["key"]: x for x in t["point_terminals"]}
    assert stats["merchant_admin"]["generated"] == 1 and stats["game_admin"]["pending_generation"] == 2 and stats["h5"]["pending"] == 1
    assert set(t["case_terminals"].values()) == {"merchant_admin"}
    # H5 没有已通过的点：按端生成 409；商户后台已全部生成也 409
    r = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={"terminal": "h5"})
    assert r.status_code == 409 and "H5" in r.json()["detail"]
    assert (await client.post(f"/api/v1/tasks/{task_id}/confirm", json={"terminal": "merchant_admin"})).status_code == 409
    # 游戏后台单独生成：增量并入，商户后台的用例与 uid 不变
    uid1 = t["result"]["cases"][0]["uid"]
    app.state.llm = StubLLM([generator_reply(make_case(case_id="TC-活动-001", module="活动运营后台-活动列表", title="活动名称必填", point_ids=[by_term["game_admin"][0]]),
                                             make_case(case_id="TC-活动-002", module="活动运营后台-活动列表", title="结束后不可编辑", point_ids=[by_term["game_admin"][1]])),
                             review_reply(True)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={"terminal": "game_admin"})
    assert resp.status_code == 200, resp.text
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert len(t["result"]["cases"]) == 3 and t["result"]["cases"][0]["uid"] == uid1 and t["points_pending_generation"] == []
    terms = t["case_terminals"]
    assert sorted(terms.values()) == ["game_admin", "game_admin", "merchant_admin"]
    # 按端导出用例：只含该端；文件名带端名；没用例的端 404；未知端 400
    r = await client.get(f"/api/v1/tasks/{task_id}/files/xlsx", params={"terminal": "merchant_admin"})
    assert r.status_code == 200 and "测试用例-商户后台.xlsx" in unquote(r.headers.get("content-disposition", ""))
    rows = _rows(r.content, tmp_path)
    assert len(rows) == 2 and rows[1][2] == "日榜降序"
    r = await client.get(f"/api/v1/tasks/{task_id}/files/xmind", params={"terminal": "game_admin"})
    assert r.status_code == 200 and len(r.content) > 100
    assert (await client.get(f"/api/v1/tasks/{task_id}/files/xlsx", params={"terminal": "h5"})).status_code == 404
    assert (await client.get(f"/api/v1/tasks/{task_id}/files/xlsx", params={"terminal": "pc"})).status_code == 400
    # 按端导出测试点：xlsx 有「端」列，只含该端；xmind 也可按端
    r = await client.get(f"/api/v1/tasks/{task_id}/points/export", params={"terminal": "h5"})
    assert r.status_code == 200 and "测试点-H5.xlsx" in unquote(r.headers.get("content-disposition", ""))
    rows = _rows(r.content, tmp_path, "p.xlsx")
    assert rows[0][2] == "端" and len(rows) == 2 and rows[1][2] == "H5" and rows[1][0] == by_term["h5"][0]
    assert (await client.get(f"/api/v1/tasks/{task_id}/points/export", params={"terminal": "game_admin", "fmt": "xmind"})).status_code == 200
    assert (await client.get(f"/api/v1/tasks/{task_id}/points/export", params={"terminal": "common"})).status_code == 404
    # 人工改端：H5 那条改到商户后台（modify），端无效 400
    h5 = by_term["h5"][0]
    r = await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [
        {"tp_id": h5, "action": "modify", "point": "未报名会员点击入口弹出报名弹窗（商户活动）", "terminal": "merchant_admin", "base_version": 1}]})
    assert r.status_code == 200, r.text
    t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert {p["tp_id"]: p.get("terminal") for m in t["analysis"]["test_points"] for p in m["points"]}[h5] == "merchant_admin"
    r = await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [{"tp_id": h5, "action": "modify", "point": "x", "terminal": "pc", "base_version": 2}]})
    assert r.status_code == 400


async def test_端到端场景_拆解前通读全文_注入各片参考_场景点单列端到端(client, monkeypatch):
    monkeypatch.setenv("TIANGONG_POINT_E2E", "true")
    get_settings.cache_clear()
    try:
        app.state.llm = StubLLM([E2E_REPLY, ANALYST_3_TERMINALS, NO_GAP])
        resp = await client.post("/api/v1/tasks", data={"text": "连赢挑战赛需求：活动运营后台、商户管理后台、H5", "confirm_points": "true"})
        assert resp.status_code == 200, resp.text
        task_id = resp.json()["task_id"]
        calls = app.state.llm.calls
        assert "端到端" in calls[0]["messages"][0]["content"] and "连赢挑战赛需求" in calls[0]["messages"][1]["content"]
        assert "端到端场景参考" in calls[1]["messages"][1]["content"] and "后台配置 → 会员报名 → 商户报表" in calls[1]["messages"][1]["content"]
        t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
        mods = {m["module"]: m["points"] for m in t["analysis"]["test_points"]}
        assert "端到端-后台配置 → 会员报名 → 商户报表" in mods and not any(k.startswith("端到端-单端") for k in mods)   # 单端的不算端到端
        e2e = mods["端到端-后台配置 → 会员报名 → 商户报表"]
        assert len(e2e) == 2 and all(p["source"] == "e2e" and p["terminal"] == "e2e" and p["status"] == "pending" and p["basis"] for p in e2e)
        assert e2e[0]["scenario"] == "后台配置 → 会员报名 → 商户报表" and e2e[0]["tp_id"].startswith("TP")
        summary = t["analysis"]["e2e"]
        assert summary["points"] == 2 and summary["scenarios"][0]["terminals"] == ["game_admin", "h5", "merchant_admin"]
        assert any(x.get("agent") == "端到端场景" and x.get("scenarios") == 1 for x in t["analysis"]["trace"])
        stats = {x["key"]: x["total"] for x in t["point_terminals"]}
        assert stats["e2e"] == 2
        # 端到端的点也可以单独生成
        for p in e2e:
            await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [{"tp_id": p["tp_id"], "action": "approve"}]})
        app.state.llm = StubLLM([generator_reply(make_case(case_id="TC-E2E-001", module=e2e[0]["module"] if "module" in e2e[0] else "端到端-后台配置 → 会员报名 → 商户报表",
                                                           title="跨端金额一致", point_ids=[e2e[0]["tp_id"], e2e[1]["tp_id"]])), review_reply(True)])
        r = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={"terminal": "e2e"})
        assert r.status_code == 200, r.text
        t = (await client.get(f"/api/v1/tasks/{task_id}")).json()
        assert list(t["case_terminals"].values()) == ["e2e"]
    finally:
        monkeypatch.delenv("TIANGONG_POINT_E2E", raising=False)
        get_settings.cache_clear()


async def test_端到端场景故障_只丢这一步_拆解照常(client, monkeypatch):
    monkeypatch.setenv("TIANGONG_POINT_E2E", "true")
    get_settings.cache_clear()
    try:
        class _FailFirst(StubLLM):
            async def chat(self, messages, model=None, require_vision=False, **overrides):
                if not self.calls:
                    self.calls.append({"messages": messages, "model": model})
                    raise RuntimeError("端到端场景模型故障")
                return await super().chat(messages, model=model, require_vision=require_vision, **overrides)

        app.state.llm = _FailFirst([ANALYST_3_TERMINALS, NO_GAP])
        resp = await client.post("/api/v1/tasks", data={"text": "需求", "confirm_points": "true"})
        assert resp.status_code == 200, resp.text
        t = (await client.get(f"/api/v1/tasks/{resp.json()['task_id']}")).json()
        assert t["analysis"]["e2e"].get("error") and sum(len(m["points"]) for m in t["analysis"]["test_points"]) == 4
        assert "端到端场景参考" not in app.state.llm.calls[1]["messages"][1]["content"]
    finally:
        monkeypatch.delenv("TIANGONG_POINT_E2E", raising=False)
        get_settings.cache_clear()


EXTRACTION_MODULES = [
    {"name": "需求说明-功能清单", "items": [{"text": "活动支持日榜与月榜", "sources": ["需求.zip › 需求说明"]}]},
    {"name": "活动运营后台 / 活动列表-基本配置", "items": [{"text": "活动名称必填", "sources": ["需求.zip › 活动运营后台 / 活动列表"]},
                                             {"text": "报名费 68 元", "sources": ["需求.zip › 活动运营后台 / 活动列表"]}]},
    {"name": "连赢活动统计-日榜统计", "items": [{"text": "按商户维度统计日榜", "sources": ["需求.zip › 连赢活动统计"]}]},
    {"name": "商户管理后台 / 日榜排名-排序", "items": [{"text": "日榜按连赢次数降序", "sources": ["需求.zip › 商户管理后台 / 日榜排名"]}]},
    {"name": "PC / 活动入口-报名", "items": [{"text": "未报名会员点击入口弹出报名弹窗", "sources": ["需求.zip › PC / 活动入口"]}]},
]
CLASSIFY_REPLY = json.dumps({"sections": [
    {"key": "连赢活动统计", "terminal": "merchant_admin", "reason": "按商户维度统计"},
    {"key": "PC", "terminal": "h5", "reason": "会员前台页面"},
    {"key": "活动运营后台", "terminal": "game_admin"},
    {"key": "不存在的章节", "terminal": "h5"}, {"key": "商户管理后台", "terminal": "bogus"},
]}, ensure_ascii=False)


async def test_章节归端模型判定_只认输入里的章节与端清单():
    from app.agents.quality import run_terminal_classify

    llm = StubLLM([CLASSIFY_REPLY])
    r = await run_terminal_classify(llm, EXTRACTION_MODULES, None)
    assert r["terminals"] == {"连赢活动统计": "merchant_admin", "PC": "h5", "活动运营后台": "game_admin"}
    assert r["reasons"]["连赢活动统计"] == "按商户维度统计"
    user = llm.calls[0]["messages"][1]["content"]
    assert "连赢活动统计" in user and "按商户维度统计日榜" in user and "merchant_admin" in user


async def test_提取完成后三端汇总_模型判定与关键词兜底_人工改端进任务(client):
    await client.post("/api/v1/projects", json={"name": "P"})
    r = (await client.post("/api/v1/requirements", data={"project": "P", "title": "连赢", "text": "连赢挑战赛需求原文"})).json()
    rid = r["req_id"]
    rstore = app.state.requirements
    rstore.set_extraction(rid, {"text": "x", "modules": EXTRACTION_MODULES, "terminals": {"连赢活动统计": "merchant_admin", "PC": "h5"}},
                          {"id": "m1", "inputs": [], "chars": 10, "modules": 5})
    r = (await client.get(f"/api/v1/requirements/{rid}")).json()
    ex = r["extraction"]
    assert ex["section_terminals"] == {"需求说明": "common", "活动运营后台": "game_admin", "连赢活动统计": "merchant_admin",
                                        "商户管理后台": "merchant_admin", "PC": "h5"}
    s = ex["terminal_summary"]
    assert [t["key"] for t in s["terminals"]] == ["h5", "game_admin", "merchant_admin", "common"]
    merchant = next(t for t in s["terminals"] if t["key"] == "merchant_admin")
    assert [x["key"] for x in merchant["sections"]] == ["连赢活动统计", "商户管理后台"] and merchant["items"] == 2
    assert s["total"]["items"] == 6 and [o["key"] for o in ex["terminal_options"]] == ["h5", "game_admin", "merchant_admin", "common", "other"]
    # 人工把「连赢活动统计」改到游戏后台；不存在的章节 404、无效端 400
    r = await client.post(f"/api/v1/requirements/{rid}/extraction/terminals", json={"section": "连赢活动统计", "terminal": "game_admin"})
    assert r.status_code == 200, r.text
    assert r.json()["extraction"]["section_terminals"]["连赢活动统计"] == "game_admin"
    assert (await client.post(f"/api/v1/requirements/{rid}/extraction/terminals", json={"section": "没有", "terminal": "h5"})).status_code == 404
    assert (await client.post(f"/api/v1/requirements/{rid}/extraction/terminals", json={"section": "PC", "terminal": "e2e"})).status_code == 400
    # 发起测试设计：章节 → 端进任务上下文，拆解出的模块按章节归端（「连赢活动统计-…」归游戏后台、「PC-…」归 H5）
    app.state.llm = StubLLM([json.dumps({"modules": [
        {"module": "连赢活动统计-日榜统计", "points": [{"point": "按商户维度统计日榜", "dimension": "数据", "ref": "a"}]},
        {"module": "PC-活动入口-报名", "points": [{"point": "弹出报名弹窗", "dimension": "正常流程", "ref": "b"}]},
        {"module": "商户管理后台-日榜排名-排序", "points": [{"point": "降序", "dimension": "数据", "ref": "c"}]},
    ], "blind_spots": []}, ensure_ascii=False), NO_GAP])
    resp = await client.post(f"/api/v1/requirements/{rid}/design", json={"async_mode": False})
    assert resp.status_code == 200, resp.text
    t = (await client.get(f"/api/v1/tasks/{resp.json()['task_id']}")).json()
    assert t["context"]["terminals"]["连赢活动统计"] == "game_admin" and t["context"]["terminals"]["PC"] == "h5"
    assert t["analysis"]["terminal_rules"]["sections"]["PC"] == "h5"
    terms = {m["module"]: {p["terminal"] for p in m["points"]} for m in t["analysis"]["test_points"]}
    assert terms == {"连赢活动统计-日榜统计": {"game_admin"}, "PC-活动入口-报名": {"h5"}, "商户管理后台-日榜排名-排序": {"merchant_admin"}}
