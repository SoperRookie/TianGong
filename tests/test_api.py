import httpx
import pytest
from asgi_lifespan import LifespanManager

from app.main import app
from tests.stubs import ANALYST_REPLY, StubLLM, generator_reply, make_case, review_reply


@pytest.fixture
async def client():
    async with LifespanManager(app):
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
            yield c


async def test_健康检查与模型列表(client):
    resp = await client.get("/health")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}

    resp = await client.get("/api/v1/models")
    assert resp.status_code == 200
    data = resp.json()
    assert data["default_model"] == "gpt-6-astra"
    names = [m["name"] for m in data["models"]]
    assert "gpt-6-astra" in names
    for m in data["models"]:
        assert "api_key_env" not in m


async def test_任务端到端_上传文本生成并下载(client):
    app.state.llm = StubLLM([ANALYST_REPLY, generator_reply(make_case()), review_reply(True)])

    resp = await client.post(
        "/api/v1/tasks",
        files={"files": ("需求.txt", "# 登录模块\n支持账号密码登录".encode(), "text/plain")},
        data={"text": "补充：连续错误 5 次锁定账号"},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "completed"
    assert data["case_count"] == 1
    assert data["passed"] is True

    # 任务查询
    resp = await client.get(f"/api/v1/tasks/{data['task_id']}")
    assert resp.status_code == 200
    assert resp.json()["sources"] == ["需求.txt", "text"]

    # 多格式同步导出下载（F-5-1/2/3/4）
    xlsx = await client.get(data["downloads"]["xlsx"])
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
    csv_resp = await client.get(data["downloads"]["csv"])
    assert csv_resp.status_code == 200
    assert csv_resp.content.startswith(b"\xef\xbb\xbf")
    xmind = await client.get(data["downloads"]["xmind"])
    assert xmind.status_code == 200
    assert xmind.content.startswith(b"PK")  # zip 魔数


async def test_解析预览接口(client):
    resp = await client.post(
        "/api/v1/parse",
        files={"files": ("需求.txt", "# 登录模块\n支持账号密码登录".encode(), "text/plain")},
        data={"text": "补充说明"},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert len(data["documents"]) == 2
    assert data["documents"][0]["source"] == "需求.txt"
    assert data["documents"][0]["sections"][0]["title"] == "登录模块"
    assert "登录模块" in data["merged_text"]
    assert data["estimated_chunks"] == 1


async def test_不支持的文件格式返回400(client):
    resp = await client.post(
        "/api/v1/tasks",
        files={"files": ("需求.exe", b"MZ\x00\x01\x90\x00", "application/octet-stream")},
    )
    assert resp.status_code == 400
    assert "不支持" in resp.json()["detail"]


async def test_超过大小上限返回413(client, monkeypatch):
    from app.config import Settings

    monkeypatch.setattr("app.api.routes.get_settings", lambda: Settings(max_upload_size_mb=0))
    resp = await client.post(
        "/api/v1/tasks",
        files={"files": ("需求.txt", b"some content", "text/plain")},
    )
    assert resp.status_code == 413


async def test_空输入返回400(client):
    resp = await client.post("/api/v1/tasks", data={"text": "  "})
    assert resp.status_code == 400


async def test_模板API_上传识别调整并按模板生成(client):
    from io import BytesIO

    from openpyxl import Workbook, load_workbook

    # 1. 上传模板 → 自动识别（F-4-1）
    wb = Workbook()
    wb.active.append(["用例ID", "所属模块", "用例名称", "优先级", "操作步骤", "预期结果", "测试类型"])
    buf = BytesIO()
    wb.save(buf)
    resp = await client.post(
        "/api/v1/templates",
        files={"file": ("团队模板.xlsx", buf.getvalue())},
        data={"name": "团队模板"},
    )
    assert resp.status_code == 200
    template = resp.json()
    assert {c["name"]: c["maps_to"] for c in template["columns"]}["测试类型"] == "custom"

    # 2. 字段映射调整（F-4-3）：给自定义列补充说明与枚举
    for col in template["columns"]:
        if col["name"] == "测试类型":
            col["description"] = "用例的测试类型"
            col["enum_values"] = ["功能", "边界"]
    resp = await client.put(f"/api/v1/templates/{template['template_id']}", json=template)
    assert resp.status_code == 200

    # 3. 模板库列表与设默认（F-4-4）
    resp = await client.post(f"/api/v1/templates/{template['template_id']}/default")
    assert resp.status_code == 200
    resp = await client.get("/api/v1/templates")
    assert resp.json()["default_id"] == template["template_id"]

    # 4. 按模板生成：导出 Excel 表头与模板 100% 一致（验收 3）
    app.state.llm = StubLLM(
        [
            ANALYST_REPLY,
            generator_reply(make_case(extras={"测试类型": "功能"})),
            review_reply(True),
        ]
    )
    resp = await client.post(
        "/api/v1/tasks",
        data={"text": "登录需求", "template_id": template["template_id"]},
    )
    assert resp.status_code == 200
    xlsx = await client.get(resp.json()["downloads"]["xlsx"])
    ws = load_workbook(BytesIO(xlsx.content)).active
    assert [c.value for c in ws[1]] == ["用例ID", "所属模块", "用例名称", "优先级", "操作步骤", "预期结果", "测试类型", "来源测试点"]

    # 清理：恢复内置默认，避免影响其他用例
    await client.post("/api/v1/templates/builtin-default/default")


async def test_拆解确认两阶段流程(client):
    # 阶段一：confirm_points=true 只做拆解
    app.state.llm = StubLLM([ANALYST_REPLY])
    resp = await client.post(
        "/api/v1/tasks",
        data={"text": "登录需求", "confirm_points": "true"},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "awaiting_confirmation"
    assert data["test_points"][0]["module"] == "登录"
    task_id = data["task_id"]

    # 阶段二：用户修改测试点后确认，继续生成并导出
    app.state.llm = StubLLM([generator_reply(make_case()), review_reply(True)])
    resp = await client.post(
        data["confirm_url"],
        json={"test_points": [{"module": "登录", "points": ["修改后的点"]}]},
    )
    assert resp.status_code == 200
    done = resp.json()
    assert done["status"] == "completed"
    assert done["case_count"] == 1
    assert "xlsx" in done["downloads"]
    # 生成收到的是修改后的测试点
    assert "修改后的点" in app.state.llm.calls[0]["messages"][1]["content"]

    # 重复确认返回 409
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})
    assert resp.status_code == 409


async def test_指定不存在的模板返回404(client):
    resp = await client.post("/api/v1/tasks", data={"text": "需求", "template_id": "ghost"})
    assert resp.status_code == 404


async def test_指定未注册模型返回400(client):
    # 使用真实 LLMClient：注册表在发起网络调用前即校验模型名
    resp = await client.post(
        "/api/v1/tasks",
        data={"text": "登录需求", "model": "不存在的模型"},
    )
    assert resp.status_code == 400
    assert "未注册" in resp.json()["detail"]


async def test_确认接口幂等_重复确认拒绝且故障可恢复重试(client):
    app.state.llm = StubLLM([ANALYST_REPLY])
    resp = await client.post(
        "/api/v1/tasks", data={"text": "登录需求", "confirm_points": "true"}
    )
    task_id = resp.json()["task_id"]

    # 生成故障（模型输出无法解析）：返回 502 且恢复待确认状态，可再次点击确认
    app.state.llm = StubLLM(["不是JSON", "还不是JSON"])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})
    assert resp.status_code == 502
    record = (await client.get(f"/api/v1/tasks/{task_id}")).json()
    assert record["status"] == "awaiting_confirmation"

    # 恢复后再次确认成功
    app.state.llm = StubLLM([generator_reply(make_case()), review_reply(True)])
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})
    assert resp.status_code == 200 and resp.json()["status"] == "completed"

    # 完成后重复确认：409（点击一次即生效，重复请求被状态锁拒绝）
    resp = await client.post(f"/api/v1/tasks/{task_id}/confirm", json={})
    assert resp.status_code == 409


async def test_任务删除_恢复_批量删除_列表与用例中心隐藏(client):
    case = make_case()
    app.state.llm = StubLLM([ANALYST_REPLY, generator_reply(case), review_reply(True)] * 3)
    ids = [(await client.post("/api/v1/tasks", data={"text": f"需求 {i}"})).json()["task_id"] for i in range(3)]
    a, b, c = ids
    assert {t["task_id"] for t in (await client.get("/api/v1/tasks")).json()["tasks"]} >= set(ids)
    for tid in ids:  # 用例中心只收审核通过的正式用例
        await client.post(f"/api/v1/tasks/{tid}/review", json={"items": [{"case_id": case["case_id"], "action": "accept"}]})
    total_cases = (await client.get("/api/v1/cases?page=0")).json()["total"]
    assert total_cases >= 3

    # 单条删除：详情 404、列表不再出现、用例中心不再计入；带 include_deleted 可见并有删除时间 / 删除人
    resp = await client.delete(f"/api/v1/tasks/{a}")
    assert resp.status_code == 200 and resp.json()["deleted"] is True and resp.json()["deleted_at"]
    assert (await client.get(f"/api/v1/tasks/{a}")).status_code == 404
    assert a not in {t["task_id"] for t in (await client.get("/api/v1/tasks")).json()["tasks"]}
    assert (await client.get("/api/v1/cases?page=0")).json()["total"] == total_cases - 1
    gone = [t for t in (await client.get("/api/v1/tasks?include_deleted=true")).json()["tasks"] if t["task_id"] == a]
    assert gone and gone[0]["deleted_at"] and gone[0]["deleted_by"] == "anonymous"
    assert (await client.delete(f"/api/v1/tasks/{a}")).status_code == 404  # 已删除的再删视同不存在
    # 恢复
    assert (await client.post(f"/api/v1/tasks/{a}/restore")).json()["deleted"] is False
    assert (await client.get(f"/api/v1/tasks/{a}")).status_code == 200
    assert (await client.get("/api/v1/cases?page=0")).json()["total"] == total_cases
    assert (await client.post(f"/api/v1/tasks/{a}/restore")).status_code == 404  # 未删除的无从恢复
    # 执行中的任务不可删
    app.state.tasks.set_progress(b, status="running")
    assert (await client.delete(f"/api/v1/tasks/{b}")).status_code == 409
    app.state.tasks.set_progress(b, status="completed")
    # 批量删除：不存在的跳过并带原因，其余生效
    r = (await client.post("/api/v1/tasks/batch-delete", json={"ids": [a, b, c, "nope", a]})).json()
    assert r["deleted"] == [a, b, c] and r["skipped"] == [{"task_id": "nope", "reason": "任务不存在: nope"}]
    assert not {a, b, c} & {t["task_id"] for t in (await client.get("/api/v1/tasks")).json()["tasks"]}
    # 操作日志
    actions = [(x["action"], x["target"], x["detail"]) for x in (await client.get("/api/v1/audit?days=0&kind=task")).json()["items"]]
    assert ("删除任务", a, "") in actions and ("恢复任务", a, "") in actions
    assert ("批量删除任务", "", "删除 3 个，跳过 1 个") in actions


async def test_修改测试点可维护设计需求页来源(client):
    import json

    reply = json.dumps({"modules": [{"module": "登录页", "points": [
        {"point": "密码错误 5 次锁定", "dimension": "异常流程", "ref": "3.2 登录失败处理 › 密码错误 5 次锁定"},
        {"point": "正常登录", "dimension": "正常流程"}]}], "blind_spots": []}, ensure_ascii=False)
    app.state.llm = StubLLM([reply])
    task_id = (await client.post("/api/v1/tasks", data={"text": "登录需求", "confirm_points": "true"})).json()["task_id"]
    pts = (await client.get(f"/api/v1/tasks/{task_id}")).json()["analysis"]["test_points"][0]["points"]
    assert pts[0]["ref"] == "3.2 登录失败处理 › 密码错误 5 次锁定" and "ref" not in pts[1]
    p1, p2 = pts[0]["tp_id"], pts[1]["tp_id"]
    # 不传 ref：保留；传 ref：改写；传空串：清空
    await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [{"tp_id": p1, "action": "modify", "point": "密码连续错误 5 次锁定账号"}]})
    await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [{"tp_id": p2, "action": "modify", "point": "正常登录", "ref": "登录页 › 账号密码登录"}]})
    pts = (await client.get(f"/api/v1/tasks/{task_id}")).json()["analysis"]["test_points"][0]["points"]
    assert pts[0]["ref"] == "3.2 登录失败处理 › 密码错误 5 次锁定" and pts[1]["ref"] == "登录页 › 账号密码登录"
    await client.post(f"/api/v1/tasks/{task_id}/points/review", json={"items": [{"tp_id": p1, "action": "modify", "point": "密码连续错误 5 次锁定账号", "ref": ""}]})
    pts = (await client.get(f"/api/v1/tasks/{task_id}")).json()["analysis"]["test_points"][0]["points"]
    assert not pts[0].get("ref")
    # 手工新增带来源
    r = (await client.post(f"/api/v1/tasks/{task_id}/points/add", json={"points": [{"module": "登录页", "point": "验证码错误提示", "ref": "登录页 › 验证码"}]})).json()
    assert r["added"][0]["ref"] == "登录页 › 验证码"
