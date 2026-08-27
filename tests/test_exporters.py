import csv

from openpyxl import load_workbook

from app.exporters import export_csv, export_excel
from app.templates import TestCase
from tests.stubs import make_case


def _cases():
    return [
        TestCase.model_validate(make_case()),
        TestCase.model_validate(
            make_case(
                case_id="TC-登录-002",
                priority="P0",
                title="验证密码错误提示",
                steps=[
                    {"action": "输入错误密码", "expected": "提示「账号或密码错误」"},
                    {"action": "连续错误 5 次", "expected": "账号锁定 30 分钟"},
                ],
            )
        ),
    ]


def test_excel导出(tmp_path):
    path = export_excel(_cases(), tmp_path / "用例.xlsx")
    ws = load_workbook(str(path)).active

    assert [c.value for c in ws[1]] == [
        "用例编号", "所属模块", "用例标题", "优先级", "前置条件", "测试步骤", "预期结果", "关键词", "备注",
    ]
    assert ws.max_row == 3
    # 步骤与预期结果编号对应
    assert ws.cell(row=3, column=6).value == "1. 输入错误密码\n2. 连续错误 5 次"
    assert ws.cell(row=3, column=7).value == "1. 提示「账号或密码错误」\n2. 账号锁定 30 分钟"
    # P0 优先级条件着色（红）
    assert ws.cell(row=3, column=4).fill.fgColor.rgb.endswith("F4CCCC")


def test_csv导出_带BOM(tmp_path):
    path = export_csv(_cases(), tmp_path / "用例.csv")

    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM

    with path.open(encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    assert rows[0][0] == "用例编号"
    assert rows[1][0] == "TC-登录-001"
    assert len(rows) == 3
