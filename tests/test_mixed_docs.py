"""图文混排文档：内嵌图片提取、位置保持、Vision 回填、小图过滤、失败不阻塞。"""

import io
import os

import fitz
from docx import Document
from PIL import Image

from app.parsers import enrich_images, parse_file
from tests.stubs import StubLLM


def _png_bytes(width=400, height=300) -> bytes:
    # 噪声图：不可压缩，保证超过小图过滤门槛（3KB）
    img = Image.frombytes("RGB", (width, height), os.urandom(width * height * 3))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _mixed_docx(path):
    doc = Document()
    doc.add_heading("充值活动需求", level=1)
    doc.add_paragraph("活动期间充值可获赠元宝。")
    doc.add_picture(io.BytesIO(_png_bytes()))  # 原型图
    doc.add_paragraph("上图为充值弹窗原型。")
    doc.save(str(path))
    return path


async def test_docx混排_图片占位与vision回填(tmp_path):
    path = _mixed_docx(tmp_path / "混排.docx")
    doc = parse_file(path)

    # 解析阶段：占位符按位置插入，图片字节已提取
    assert len(doc.embedded_images) == 1
    assert "[[图片:1]]" in doc.full_text
    text = doc.full_text
    assert text.index("活动期间充值") < text.index("[[图片:1]]") < text.index("上图为充值弹窗原型")

    # 回填阶段：Vision 理解结果替换占位符
    llm = StubLLM(["- 弹窗含充值档位按钮：100/500/1000 元"])
    enriched = await enrich_images(doc, llm)

    assert llm.calls[0]["require_vision"] is True
    full = enriched.full_text
    assert "[[图片:1]]" not in full or "文档内图片理解" in full
    assert "充值档位按钮" in full
    assert full.index("活动期间充值") < full.index("充值档位按钮") < full.index("上图为充值弹窗原型")
    assert enriched.embedded_images == []


async def test_装饰性小图跳过vision(tmp_path):
    doc = Document()
    doc.add_paragraph("正文")
    doc.add_picture(io.BytesIO(_png_bytes(width=40, height=40)))  # 小图标
    p = tmp_path / "小图.docx"
    doc.save(str(p))

    parsed = parse_file(p)
    llm = StubLLM([])  # 不应产生任何调用
    enriched = await enrich_images(parsed, llm)

    assert llm.calls == []
    assert "装饰性小图" in enriched.full_text


async def test_单图失败不阻塞整体(tmp_path):
    path = _mixed_docx(tmp_path / "失败.docx")
    doc = parse_file(path)

    class FailingLLM:
        async def chat(self, *a, **k):
            raise RuntimeError("Vision 服务不可用")

    enriched = await enrich_images(doc, FailingLLM())
    assert "解析失败" in enriched.full_text
    assert "上图为充值弹窗原型" in enriched.full_text  # 文本内容不受影响


def test_pdf混排_按页提取图片(tmp_path):
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Recharge Requirement", fontsize=20)
    page.insert_text((72, 110), "Detail rules as below.", fontsize=11)
    page.insert_image(fitz.Rect(72, 140, 372, 365), stream=_png_bytes())
    path = tmp_path / "混排.pdf"
    pdf.save(str(path))
    pdf.close()

    doc = parse_file(path)

    assert len(doc.embedded_images) == 1
    assert "[[图片:1]]" in doc.full_text
    assert "Detail rules as below." in doc.full_text
