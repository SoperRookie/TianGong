"""自定义模板：识别 / 校验 / 模板库 / 按模板导出。"""

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from app.agents.graph import rule_check
from app.exporters import export_excel
from app.templates import (
    CustomTemplate,
    TemplateField,
    TemplateParseError,
    TemplateStore,
    TestCase,
    recognize_template,
)
from tests.stubs import make_case


def _team_template_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.append(["用例ID", "所属模块", "用例名称", "优先级", "前置条件", "操作步骤", "预期结果", "测试结果"])
    ws.append(["TC-001", "登录", "验证登录", "P0", "已注册", "1.输入密码", "1.登录成功", "通过"])
    ws.append(["TC-002", "登录", "验证登出", "P1", "已登录", "1.点击登出", "1.回到首页", "不通过"])
    wb.save(str(path))
    return path


def test_excel模板字段自动识别(tmp_path):
    path = _team_template_xlsx(tmp_path / "团队模板.xlsx")
    template = recognize_template(path)

    assert template.name == "团队模板"
    mapping = {c.name: c.maps_to for c in template.columns}
    assert mapping == {
        "用例ID": "case_id",
        "所属模块": "module",
        "用例名称": "title",
        "优先级": "priority",
        "前置条件": "precondition",
        "操作步骤": "steps",
        "预期结果": "expected",
        "测试结果": "custom",  # 无法映射的列识别为自定义字段
    }
    priority = template.column_for("priority")
    assert priority.enum_values == ["P0", "P1"]  # 从数据行提取枚举
    assert priority.required is True


def test_word表格模板识别(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for i, h in enumerate(["编号", "标题", "优先级", "步骤"]):
        table.cell(0, i).text = h
    path = tmp_path / "模板.docx"
    doc.save(str(path))

    template = recognize_template(path)
    assert [c.maps_to for c in template.columns] == ["case_id", "title", "priority", "steps"]


def test_无表格word模板报错(tmp_path):
    doc = Document()
    doc.add_paragraph("只有正文")
    path = tmp_path / "空模板.docx"
    doc.save(str(path))
    with pytest.raises(TemplateParseError, match="表格"):
        recognize_template(path)


def _custom_template():
    return CustomTemplate(
        name="含自定义字段模板",
        columns=[
            TemplateField(name="用例编号", maps_to="case_id", required=True),
            TemplateField(name="模块", maps_to="module", required=True),
            TemplateField(name="标题", maps_to="title", required=True),
            TemplateField(name="优先级", maps_to="priority", enum_values=["P0", "P1"]),
            TemplateField(name="测试步骤", maps_to="steps", required=True),
            TemplateField(name="预期结果", maps_to="expected", required=True),
            TemplateField(name="测试类型", maps_to="custom", required=True, enum_values=["功能", "边界"]),
        ],
    )


def test_模板约束校验():
    template = _custom_template()
    ok = make_case(priority="P1", extras={"测试类型": "功能"})
    missing_extra = make_case(case_id="TC-登录-002")  # 缺必填自定义字段
    bad_priority = make_case(case_id="TC-登录-003", priority="P2")  # 超出模板枚举
    bad_enum = make_case(case_id="TC-登录-004", extras={"测试类型": "性能"})

    problems = " ".join(i["problem"] for i in rule_check([ok, missing_extra, bad_priority, bad_enum], template))
    assert "测试类型」缺失" in problems
    assert "P2 不在模板允许范围" in problems
    assert "不在枚举" in problems
    assert rule_check([ok], template) == []


def test_prompt_spec包含自定义字段说明():
    spec = _custom_template().prompt_spec()
    assert "extras[\"测试类型\"]" in spec
    assert "仅允许 P0/P1" in spec
    assert "['功能', '边界']" in spec


def test_按模板导出excel列序一致(tmp_path):
    template = _custom_template()
    case = TestCase.model_validate(make_case(priority="P0", extras={"测试类型": "功能"}))
    path = export_excel([case], tmp_path / "自定义.xlsx", template)

    ws = load_workbook(str(path)).active
    assert [c.value for c in ws[1]] == ["用例编号", "模块", "标题", "优先级", "测试步骤", "预期结果", "测试类型"]
    assert ws.cell(row=2, column=7).value == "功能"  # 自定义列取 extras


def test_模板库管理(tmp_path):
    store = TemplateStore(tmp_path / "templates.json")
    assert store.get().template_id == "builtin-default"  # 内置默认自动就位

    template = store.save(_custom_template())
    store.set_default(template.template_id)
    assert store.get().name == "含自定义字段模板"

    # 持久化：新实例可读回
    store2 = TemplateStore(tmp_path / "templates.json")
    assert store2.default_id == template.template_id
    assert store2.get(template.template_id).columns[6].name == "测试类型"

    # 删除默认模板后回落内置
    store2.delete(template.template_id)
    assert store2.get().template_id == "builtin-default"
    with pytest.raises(ValueError, match="不可删除"):
        store2.delete("builtin-default")
