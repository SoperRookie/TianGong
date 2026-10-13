"""M4b：人工用例集、人工新增/复制（草稿 → 提交评审）、Excel 导入五步校验、批量维护与乐观锁。"""

import httpx
import pytest
from asgi_lifespan import LifespanManager

from app.main import app
from tests.stubs import make_case


@pytest.fixture
async def client():
    async with LifespanManager(app):
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
            yield c


async def _manual_task(client) -> str:
    await client.post("/api/v1/projects", json={"name": "P"})
    resp = await client.post("/api/v1/projects/P/manual-task", json={"title": "登录模块手工用例"})
    assert resp.status_code == 200, resp.text
    return resp.json()["task_id"]


async def test_人工新增_复制_草稿提交评审(client):
    tid = await _manual_task(client)
    body = {"module": "登录", "title": "验证码过期", "priority": "P1", "steps": [{"action": "等待 5 分钟后提交", "expected": "提示验证码已过期"}], "keywords": "登录、验证码"}
    r = (await client.post(f"/api/v1/tasks/{tid}/cases", json=body)).json()
    c = r["case"]
    assert c["case_id"] == "TC-登录-001" and c["source"] == "manual" and c["version"] == 1
    task = (await client.get(f"/api/v1/tasks/{tid}")).json()
    assert task["case_reviews"][c["uid"]]["status"] == "draft"
    # 草稿不能直接通过；提交评审后可通过
    resp = await client.post(f"/api/v1/tasks/{tid}/review", json={"items": [{"case_id": c["case_id"], "action": "approve"}]})
    assert resp.status_code == 409
    await client.post(f"/api/v1/tasks/{tid}/review", json={"items": [{"case_id": c["case_id"], "action": "submit"}]})
    task = (await client.get(f"/api/v1/tasks/{tid}")).json()
    assert task["case_reviews"][c["uid"]]["status"] == "pending"
    assert (await client.post(f"/api/v1/tasks/{tid}/review", json={"items": [{"case_id": c["case_id"], "action": "approve"}]})).status_code == 200
    # 复制：草稿、来源 copy、标题带副本、版本历史首版
    r = (await client.post(f"/api/v1/tasks/{tid}/cases/{c['uid']}/copy")).json()
    dup = r["case"]
    assert dup["title"].endswith("（副本）") and dup["source"] == "copy" and dup["case_id"] == "TC-登录-002"
    task = (await client.get(f"/api/v1/tasks/{tid}")).json()
    assert task["case_reviews"][dup["uid"]]["status"] == "draft"
    vers = (await client.get(f"/api/v1/tasks/{tid}/versions", params={"kind": "case", "entity_id": dup["uid"]})).json()
    assert vers["versions"][0]["source"] == "manual"
    # 非法用例被拒
    assert (await client.post(f"/api/v1/tasks/{tid}/cases", json={**body, "title": ""})).status_code == 400
    # 用例中心只收录正式用例：草稿不进入
    lib = (await client.get("/api/v1/cases", params={"project": "P"})).json()
    assert {x["review"] for x in lib["cases"]} == {"approved"} and all(x["uid"] != dup["uid"] for x in lib["cases"])


async def test_Excel导入五步校验(client, tmp_path):
    tid = await _manual_task(client)
    await client.post(f"/api/v1/tasks/{tid}/cases", json={"module": "登录", "title": "已存在的用例", "steps": [{"action": "a", "expected": "b"}]})
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["用例编号", "模块", "用例标题", "优先级", "前置条件", "操作步骤", "预期结果", "关键词", "备注"])
    ws.append(["TC-登录-001", "登录", "正常登录", "P1", "", "1. 输入正确账号", "1. 登录成功", "登录", ""])
    ws.append(["TC-登录-002", "登录", "已存在的用例", "P1", "", "1. x", "1. y", "", ""])
    ws.append(["TC-登录-003", "登录", "正常登录", "P2", "", "1. 重复标题", "1. z", "", ""])
    ws.append(["TC-登录-004", "", "缺模块缺预期", "P3", "", "1. 步骤", "", "", ""])
    ws.append(["TC-登录-005", "", "缺模块", "P3", "", "1. 步骤", "1. 结果", "", ""])
    path = tmp_path / "导入.xlsx"; wb.save(path)
    with open(path, "rb") as fh:
        resp = await client.post(f"/api/v1/tasks/{tid}/cases/import/preview", files={"file": ("导入.xlsx", fh.read(), "application/octet-stream")})
    assert resp.status_code == 200, resp.text
    pv = resp.json()
    assert pv["total"] == 5 and pv["errors"] == 2
    rows = {r["row"]: r for r in pv["rows"]}
    assert rows[1]["errors"] == []
    assert "与任务内已有用例标题重复" in rows[2]["errors"]
    assert any("与第 1 行标题重复" in e for e in rows[3]["errors"])
    assert rows[4]["errors"] == [] and any("无预期结果" in w for w in rows[4]["warnings"])  # 历史简化用例：缺预期只警告
    assert rows[5]["errors"] == [] and rows[5]["case"]["module"] == "未分组" and "缺少模块" in " ".join(rows[5]["warnings"])
    # 有错误行：不显式选择只导有效行则拒绝
    resp = await client.post(f"/api/v1/tasks/{tid}/cases/import/confirm", json={"token": pv["token"]})
    assert resp.status_code == 400 and "2 行" in resp.json()["detail"]
    resp = await client.post(f"/api/v1/tasks/{tid}/cases/import/confirm", json={"token": pv["token"], "only_valid": True})
    assert resp.status_code == 200 and resp.json()["imported"] == 3 and resp.json()["skipped"] == 2
    task = (await client.get(f"/api/v1/tasks/{tid}")).json()
    imported = [c for c in task["result"]["cases"] if c["source"] == "import"]
    assert len(imported) == 3 and all(task["case_reviews"][c["uid"]]["status"] == "draft" for c in imported)
    blank = next(c for c in imported if c["title"] == "缺模块缺预期")
    assert blank["steps"][0]["action"] == "步骤" and blank["steps"][0]["expected"] == ""
    # 令牌一次性
    assert (await client.post(f"/api/v1/tasks/{tid}/cases/import/confirm", json={"token": pv["token"], "only_valid": True})).status_code == 404
    # 模板下载
    resp = await client.get("/api/v1/cases/import-template")
    assert resp.status_code == 200 and resp.content[:2] == b"PK"


async def test_批量维护_乐观锁_回收站(client):
    tid = await _manual_task(client)
    uids = []
    for i in range(3):
        r = (await client.post(f"/api/v1/tasks/{tid}/cases", json={"module": "登录", "title": f"用例{i}", "steps": [{"action": "a", "expected": "b"}]})).json()
        uids.append(r["case"]["uid"])
    # 批量提交评审 → 全部 pending
    r = (await client.post(f"/api/v1/tasks/{tid}/cases/batch", json={"uids": uids, "action": "submit"})).json()
    assert len(r["applied"]) == 3 and all(v["status"] == "pending" for v in r["case_reviews"].values())
    # 批量改优先级：带页面版本，其中一条版本过期 → 冲突跳过
    r = (await client.post(f"/api/v1/tasks/{tid}/cases/batch",
                           json={"uids": uids, "action": "set_priority", "value": "p0", "base_versions": {uids[0]: 1, uids[1]: 99}})).json()
    assert [c["uid"] for c in r["conflicts"]] == [uids[1]] and len(r["applied"]) == 2
    by = {c["uid"]: c for c in r["cases"]}
    assert by[uids[0]]["priority"] == "P0" and by[uids[0]]["version"] == 2 and by[uids[1]]["priority"] == "P2"
    # 批量加标签（追加去重）与改模块
    r = (await client.post(f"/api/v1/tasks/{tid}/cases/batch", json={"uids": uids[:1], "action": "add_keywords", "value": "回归, 冒烟"})).json()
    assert r["cases"][0]["keywords"] == "回归、冒烟"
    r = (await client.post(f"/api/v1/tasks/{tid}/cases/batch", json={"uids": uids, "action": "set_module", "value": "账号"})).json()
    assert all(c["module"] == "账号" and c["case_id"].startswith("TC-账号-") for c in r["cases"])
    # 已通过锁定的用例跳过
    await client.post(f"/api/v1/tasks/{tid}/review", json={"items": [{"case_id": r["cases"][0]["case_id"], "action": "approve"}]})
    r = (await client.post(f"/api/v1/tasks/{tid}/cases/batch", json={"uids": uids[:1], "action": "set_priority", "value": "P3"})).json()
    assert r["skipped"][0]["reason"].startswith("已通过并锁定")
    # 批量删除 → 回收站可恢复
    r = (await client.post(f"/api/v1/tasks/{tid}/cases/batch", json={"uids": uids[1:], "action": "delete"})).json()
    assert len(r["applied"]) == 2 and len(r["cases"]) == 1
    bin_items = (await client.get(f"/api/v1/tasks/{tid}/recycle-bin")).json()["items"]
    assert len(bin_items) == 2
    assert (await client.post(f"/api/v1/tasks/{tid}/recycle-bin/restore", json={"item_id": bin_items[0]["id"]})).status_code == 200
    assert (await client.post(f"/api/v1/tasks/{tid}/cases/batch", json={"uids": uids, "action": "explode"})).status_code == 400


async def test_Excel导入多文件合并预览(client, tmp_path):
    from openpyxl import Workbook

    tid = await _manual_task(client)
    header = ["用例编号", "模块", "用例标题", "优先级", "前置条件", "操作步骤", "预期结果", "关键词", "备注"]
    wb = Workbook(); ws = wb.active; ws.append(header)
    ws.append(["TC-登录-001", "登录", "正常登录", "P1", "", "1. 输入正确账号", "1. 登录成功", "", ""])
    ws.append(["TC-登录-002", "登录", "密码错误", "P2", "", "1. 输入错误密码", "1. 提示错误", "", ""])
    a = tmp_path / "登录.xlsx"; wb.save(a)
    wb = Workbook(); ws = wb.active; ws.append(header)
    ws.append(["TC-注册-001", "注册", "正常注册", "P1", "", "1. 填写资料", "1. 注册成功", "", ""])
    ws.append(["TC-注册-002", "注册", "正常登录", "P2", "", "1. 跨文件重复标题", "1. z", "", ""])
    b = tmp_path / "注册.xlsx"; wb.save(b)
    files = [("files", ("登录.xlsx", a.read_bytes(), "application/octet-stream")),
             ("files", ("注册.xlsx", b.read_bytes(), "application/octet-stream"))]
    resp = await client.post(f"/api/v1/tasks/{tid}/cases/import/preview", files=files)
    assert resp.status_code == 200, resp.text
    pv = resp.json()
    assert pv["files"] == ["登录.xlsx", "注册.xlsx"] and pv["total"] == 4 and pv["errors"] == 1
    rows = {(r["file"], r["row"]): r for r in pv["rows"]}
    assert rows[("登录.xlsx", 1)]["errors"] == [] and rows[("注册.xlsx", 1)]["errors"] == []
    assert any("登录.xlsx 第 1 行标题重复" in e for e in rows[("注册.xlsx", 2)]["errors"])
    resp = await client.post(f"/api/v1/tasks/{tid}/cases/import/confirm", json={"token": pv["token"], "only_valid": True})
    assert resp.status_code == 200 and resp.json()["imported"] == 3 and resp.json()["skipped"] == 1
    # 一个文件都没选 → 400
    assert (await client.post(f"/api/v1/tasks/{tid}/cases/import/preview")).status_code == 400


async def test_用例中心筛选_执行状态_创建人_关键词覆盖步骤(client):
    tid = await _manual_task(client)
    mk = lambda module, title, action, **kw: {"module": module, "title": title, "priority": kw.get("priority", "P1"),
                                               "steps": [{"action": action, "expected": kw.get("expected", "成功")}], **{k: v for k, v in kw.items() if k in ("keywords", "remark", "precondition")}}
    cases = []
    for body in (mk("登录", "正常登录", "输入正确账号密码", keywords="登录"),
                 mk("登录", "锁定", "连续错误 5 次", priority="P0", remark="注意锁定时长"),
                 mk("注册", "注册成功", "填写手机号与验证码", precondition="手机号未注册")):
        c = (await client.post(f"/api/v1/tasks/{tid}/cases", json=body)).json()["case"]
        await client.post(f"/api/v1/tasks/{tid}/review", json={"items": [{"case_id": c["case_id"], "action": "submit"}]})
        await client.post(f"/api/v1/tasks/{tid}/review", json={"items": [{"case_id": c["case_id"], "action": "approve"}]})
        cases.append(c)
    lib = (await client.get("/api/v1/cases", params={"project": "P"})).json()
    assert lib["total"] == 3 and lib["module_counts"] == {"登录": 2, "注册": 1} and lib["creators"] == ["anonymous"]
    # 关键词覆盖步骤 / 备注 / 前置条件
    assert (await client.get("/api/v1/cases", params={"project": "P", "keyword": "验证码"})).json()["total"] == 1
    assert (await client.get("/api/v1/cases", params={"project": "P", "keyword": "锁定时长"})).json()["total"] == 1
    assert (await client.get("/api/v1/cases", params={"project": "P", "keyword": "未注册"})).json()["total"] == 1
    assert (await client.get("/api/v1/cases", params={"project": "P", "keyword": "不存在"})).json()["total"] == 0
    # 模块筛选不影响 module_counts（下拉始终显示全量条数）
    r = (await client.get("/api/v1/cases", params={"project": "P", "module": "注册"})).json()
    assert r["total"] == 1 and r["module_counts"] == {"登录": 2, "注册": 1}
    # 执行状态：未执行 3 条；在计划里执行一条后按状态筛
    assert (await client.get("/api/v1/cases", params={"project": "P", "exec_status": "none"})).json()["total"] == 3
    assert (await client.get("/api/v1/cases", params={"project": "P", "exec_status": "fail"})).json()["total"] == 0
    plan = (await client.post("/api/v1/plans", json={"project": "P", "name": "冒烟"})).json()
    pid = plan["plan_id"]
    assert (await client.post(f"/api/v1/plans/{pid}/cases", json={"task_id": tid})).status_code == 200
    run = (await client.post(f"/api/v1/plans/{pid}/runs", json={"name": "第 1 轮"})).json()
    items = (await client.get(f"/api/v1/plans/{pid}")).json()["items"]
    item_id = next(i["item_id"] for i in items if i["uid"] == cases[1]["uid"])
    resp = await client.post(f"/api/v1/plans/{pid}/runs/{run['run_id']}/results",
                             json={"items": [{"item_id": item_id, "status": "fail", "note": "锁定未生效", "reason": "系统缺陷"}]})
    assert resp.status_code == 200, resp.text
    assert [x["title"] for x in (await client.get("/api/v1/cases", params={"project": "P", "exec_status": "fail"})).json()["cases"]] == ["锁定"]
    assert (await client.get("/api/v1/cases", params={"project": "P", "exec_status": "none"})).json()["total"] == 2
    # 创建人
    assert (await client.get("/api/v1/cases", params={"project": "P", "created_by": "anonymous"})).json()["total"] == 3
    assert (await client.get("/api/v1/cases", params={"project": "P", "created_by": "nobody"})).json()["total"] == 0
