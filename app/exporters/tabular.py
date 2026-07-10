"""表格导出（F-5-2 Excel / F-5-3 CSV）。

列结构由模板决定（F-4-5：导出字段与模板 100% 一致），缺省用内置默认模板。
Excel：表头样式、列宽自适应、优先级条件着色。
CSV：UTF-8 with BOM，便于导入禅道 / TestLink / Jira。
"""

import csv
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.templates import CustomTemplate, TestCase, builtin_default_template

_PRIORITY_FILLS = {
    "P0": PatternFill("solid", fgColor="F4CCCC"),  # 红
    "P1": PatternFill("solid", fgColor="FCE5CD"),  # 橙
    "P2": PatternFill("solid", fgColor="FFF2CC"),  # 黄
    "P3": PatternFill("solid", fgColor="EFEFEF"),  # 灰
}
_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_MAX_COL_WIDTH = 60

_CANONICAL_GETTERS: dict[str, Callable[[TestCase], str]] = {
    "case_id": lambda c: c.case_id,
    "module": lambda c: c.module,
    "title": lambda c: c.title,
    "priority": lambda c: c.priority,
    "precondition": lambda c: c.precondition,
    "steps": lambda c: "\n".join(f"{i}. {s.action}" for i, s in enumerate(c.steps, 1)),
    "expected": lambda c: "\n".join(f"{i}. {s.expected}" for i, s in enumerate(c.steps, 1)),
    "remark": lambda c: c.remark,
}


def _headers(template: CustomTemplate) -> list[str]:
    return [col.name for col in template.columns]


def _rows(cases: list[TestCase], template: CustomTemplate) -> list[list[str]]:
    rows = []
    for case in cases:
        row = []
        for col in template.columns:
            getter = _CANONICAL_GETTERS.get(col.maps_to)
            row.append(getter(case) if getter else case.extras.get(col.name, ""))
        rows.append(row)
    return rows


def export_excel(
    cases: list[TestCase], path: str | Path, template: CustomTemplate | None = None
) -> Path:
    template = template or builtin_default_template()
    headers = _headers(template)
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in _rows(cases, template):
        ws.append(row)

    priority_col = next(
        (i + 1 for i, col in enumerate(template.columns) if col.maps_to == "priority"), None
    )
    for row_idx in range(2, ws.max_row + 1):
        if priority_col:
            cell = ws.cell(row=row_idx, column=priority_col)
            fill = _PRIORITY_FILLS.get(str(cell.value))
            if fill:
                cell.fill = fill
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    # 列宽自适应：按内容最长行估算（中文按 2 个宽度计），设上限
    for col_idx in range(1, len(headers) + 1):
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value or ""
            for line in str(value).splitlines():
                width = sum(2 if ord(ch) > 127 else 1 for ch in line)
                max_len = max(max_len, width)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, _MAX_COL_WIDTH)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def export_csv(
    cases: list[TestCase], path: str | Path, template: CustomTemplate | None = None
) -> Path:
    template = template or builtin_default_template()
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(_headers(template))
        writer.writerows(_rows(cases, template))
    return path
