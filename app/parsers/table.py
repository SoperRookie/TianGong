"""表格类需求文件（Excel / CSV）：按工作表渲染为可读文本，并提取内嵌截图交给 Vision。

策划 / 产品常把需求写在 Excel 里：一列长文字 + 旁边贴截图 + 大量空单元格。按「行 | 列」原样转出会充斥空格分隔符、
丢掉截图上下文，模型读不出结构。这里按工作表自动判别：
- 文档型（多数行只有一两个非空单元格）：逐行输出非空单元格文本，保留编号与段落；
- 数据型（有表头、多列填充）：每行按「列名：值」展开，自描述，不依赖表头位置；
内嵌图片按锚点行插入 [[图片:N]] 占位，由 enrich_images 经 Vision 理解后回填。
"""

import csv
import io
from pathlib import Path

from app.parsers.base import EmbeddedImage, ParsedDocument, Section

_IMG_MIME = {"png": "image/png", "jpeg": "image/jpeg", "jpg": "image/jpeg", "gif": "image/gif", "bmp": "image/bmp",
             "tiff": "image/tiff", "webp": "image/webp"}


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _is_data_table(rows: list[list[str]]) -> bool:
    """首行像表头（≥3 个非空、都较短）且多数行填充列数接近表头 → 数据型。"""
    if len(rows) < 2:
        return False
    header = [c for c in rows[0] if c]
    if len(header) < 3 or any(len(h) > 30 or "\n" in h for h in header):
        return False
    filled = [sum(1 for c in r if c) for r in rows[1:]]
    return sum(1 for n in filled if n >= max(2, len(header) // 2)) >= 0.6 * len(filled)


def render_rows(rows: list[list[str]], images_at: dict[int, list[str]] | None = None) -> list[str]:
    """二维表 → 文本行；images_at: 行号 → 该行后要插入的占位符。"""
    images_at = images_at or {}
    out: list[str] = []
    rows = [[_cell(c) for c in r] for r in rows]
    if _is_data_table(rows):
        header = rows[0]
        for i, r in enumerate(rows[1:], start=1):
            pairs = [f"{header[j] or f'列{j + 1}'}：{v}" for j, v in enumerate(r) if v]
            if pairs:
                out.append("；".join(pairs))
            out.extend(images_at.get(i, []))
        return out
    for i, r in enumerate(rows):
        cells = [c for c in r if c]
        if cells:
            out.append(" ｜ ".join(cells) if len(cells) > 1 else cells[0])
        out.extend(images_at.get(i, []))
    return out


class TableParser:
    suffixes = (".xlsx", ".xlsm", ".csv", ".tsv")

    def parse(self, path: Path) -> ParsedDocument:
        suffix = path.suffix.lower()
        sections: list[Section] = []
        images: list[EmbeddedImage] = []
        if suffix in (".csv", ".tsv"):
            from app.parsers.base import read_text_any

            rows = list(csv.reader(io.StringIO(read_text_any(path)), delimiter="\t" if suffix == ".tsv" else ","))
            lines = render_rows(rows)
            if not lines:
                raise ValueError(f"{path.name} 没有可读取的表格内容")
            sections.append(Section(level=0, content="\n".join(lines)))
            return ParsedDocument(source=path.name, doc_type=suffix.lstrip("."), sections=sections)

        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
            images_at: dict[int, list[str]] = {}
            for im in getattr(ws, "_images", []) or []:
                try:
                    data = im._data()
                    fmt = (getattr(im, "format", "") or "png").lower()
                    anchor = getattr(im, "anchor", None)
                    row = getattr(getattr(anchor, "_from", None), "row", None)
                    row = int(row) if row is not None else len(rows) - 1
                except Exception:
                    continue
                placeholder = f"[[图片:{len(images) + 1}]]"
                images.append(EmbeddedImage(placeholder=placeholder, data=data, mime=_IMG_MIME.get(fmt, "image/png")))
                images_at.setdefault(min(max(row, 0), max(len(rows) - 1, 0)), []).append(placeholder)
            lines = render_rows(rows, images_at)
            if not lines:
                continue
            sections.append(Section(level=1, title=ws.title))
            # 占位符单独成段，便于 enrich_images 原位回填
            buf: list[str] = []
            for line in lines:
                if line.startswith("[[图片:"):
                    if buf:
                        sections.append(Section(level=0, content="\n".join(buf)))
                        buf = []
                    sections.append(Section(level=0, content=line))
                else:
                    buf.append(line)
            if buf:
                sections.append(Section(level=0, content="\n".join(buf)))
        if not sections:
            raise ValueError(f"{path.name} 没有可读取的表格内容")
        return ParsedDocument(source=path.name, doc_type=suffix.lstrip("."), sections=sections, embedded_images=images)
