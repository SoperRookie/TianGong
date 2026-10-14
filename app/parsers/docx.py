"""Word 解析（F-2-2）：提取正文、标题层级、表格、列表，保留章节结构。

图文混排支持：段落中的内嵌图片按位置生成占位 Section（[[图片:N]]），
图片字节存入 embedded_images，由 enrich_images 经 Vision 理解后原位回填。
"""

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from app.parsers.base import EmbeddedImage, ParsedDocument, Section

# 兼容英文样式 "Heading 1" 与中文样式 "标题 1"
_HEADING_STYLE_RE = re.compile(r"^(?:heading|标题)\s*(\d)", re.IGNORECASE)
_LIST_STYLE_RE = re.compile(r"list|列表|项目符号|编号", re.IGNORECASE)


class DocxParser:
    suffixes = (".docx",)

    def parse(self, path: Path) -> ParsedDocument:
        document = Document(str(path))
        sections: list[Section] = []
        images: list[EmbeddedImage] = []
        buffer: list[str] = []

        def flush() -> None:
            content = "\n".join(buffer).strip()
            if content:
                sections.append(Section(level=0, content=content))
            buffer.clear()

        def add_image(data: bytes, mime: str) -> None:
            flush()
            placeholder = f"[[图片:{len(images) + 1}]]"
            sections.append(Section(level=0, content=placeholder))
            images.append(EmbeddedImage(placeholder=placeholder, data=data, mime=mime))

        for kind, obj in self._iter_body(document):
            if kind == "table":
                # 表格文字仍统一收进 tables；单元格里的图片（Axure / 产品常把页面截图放在表格里）按位置提取，
                # 合并单元格会重复返回同一单元格，按关系 id 去重
                seen: set[str] = set()
                for row in obj.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            for rid, data, mime in self._para_images(para, document):
                                if rid not in seen:
                                    seen.add(rid)
                                    add_image(data, mime)
                continue
            para = obj
            for _, data, mime in self._para_images(para, document):
                add_image(data, mime)
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name if para.style else ""
            heading = _HEADING_STYLE_RE.match(style_name)
            if heading:
                flush()
                sections.append(Section(level=int(heading.group(1)), title=text))
            elif _LIST_STYLE_RE.search(style_name):
                buffer.append(f"- {text}")
            else:
                buffer.append(text)
        flush()

        tables = [
            [[cell.text.strip() for cell in row.cells] for row in table.rows]
            for table in document.tables
        ]
        return ParsedDocument(
            source=path.name, doc_type="docx", sections=sections, tables=tables,
            embedded_images=images,
        )

    @staticmethod
    def _iter_body(document):
        """按正文顺序产出 ("para", Paragraph) / ("table", Table)：python-docx 的 document.paragraphs 不含表格里的段落。"""
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                yield "para", Paragraph(child, document)
            elif child.tag == qn("w:tbl"):
                yield "table", Table(child, document)

    @staticmethod
    def _para_images(para, document) -> list[tuple[str, bytes, str]]:
        """提取段落内嵌图片 (关系 id, 字节, mime)（w:drawing → a:blip 的关系引用）。"""
        results: list[tuple[str, bytes, str]] = []
        for blip in para._element.findall(".//" + qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if not rid:
                continue
            part = document.part.related_parts.get(rid)
            if part is not None:
                results.append((rid, part.blob, part.content_type))
        return results
