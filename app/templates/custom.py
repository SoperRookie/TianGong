"""自定义模板（F-4-1/3）：上传 Excel/Word 模板，自动识别字段结构。

内部用例数据始终使用规范字段（TestCase），模板列通过 maps_to 映射到规范字段；
无法映射的列作为自定义字段，生成值存放在 TestCase.extras[列名]。
优先级内部统一 P0-P3（2026-07-10 决策）；模板枚举若为其子集则按模板枚举校验。
"""

import re
import uuid
from pathlib import Path
from typing import Literal

from docx import Document
from openpyxl import load_workbook
from pydantic import BaseModel, Field

CanonicalField = Literal[
    "case_id", "module", "title", "priority", "precondition", "steps", "expected",
    "keywords", "remark", "custom",
]

# 常见中文/英文列名 → 规范字段 的启发式映射（按顺序优先匹配）
_CANONICAL_PATTERNS: list[tuple[CanonicalField, list[str]]] = [
    ("case_id", ["用例编号", "用例id", "caseid", "case id", "编号"]),
    ("module", ["所属模块", "功能模块", "模块"]),
    ("title", ["用例标题", "用例名称", "标题", "名称"]),
    ("priority", ["优先级", "级别", "等级"]),
    ("precondition", ["前置条件", "预置条件", "前提条件", "前提"]),
    ("expected", ["预期结果", "期望结果", "预期输出", "预期"]),
    ("steps", ["测试步骤", "操作步骤", "执行步骤", "步骤"]),
    ("keywords", ["关键词", "标签", "keywords", "tags"]),
    ("remark", ["备注", "说明", "注释"]),
]

PRIORITY_ENUM = ["P0", "P1", "P2", "P3"]


class TemplateField(BaseModel):
    name: str = Field(description="模板中的列名（导出表头原样使用）")
    maps_to: CanonicalField = Field(default="custom", description="映射到的规范字段")
    required: bool = False
    enum_values: list[str] | None = Field(default=None, description="取值枚举（如优先级）")
    description: str = Field(default="", description="字段含义说明，注入生成 Prompt（F-4-3 可调整）")


class CustomTemplate(BaseModel):
    template_id: str = Field(default_factory=lambda: uuid.uuid4().hex[:12])
    name: str
    columns: list[TemplateField]

    def column_for(self, canonical: str) -> TemplateField | None:
        return next((c for c in self.columns if c.maps_to == canonical), None)

    def custom_columns(self) -> list[TemplateField]:
        return [c for c in self.columns if c.maps_to == "custom"]

    def priority_enum(self) -> list[str]:
        """模板枚举为 P0-P3 子集时生效，否则回落内部标准（含团队 p4 合并 P3 的决策）。"""
        col = self.column_for("priority")
        if col and col.enum_values:
            normalized = [v.strip().upper() for v in col.enum_values]
            if normalized and set(normalized) <= set(PRIORITY_ENUM):
                return normalized
        return PRIORITY_ENUM

    def prompt_spec(self) -> str:
        from app.templates.default import DEFAULT_TEMPLATE

        canonical_specs = DEFAULT_TEMPLATE.field_specs
        lines: list[str] = []
        for col in self.columns:
            if col.maps_to == "priority":
                spec = f"优先级，仅允许 {'/'.join(self.priority_enum())}：P0=冒烟与核心链路；P1=主功能正常流与重要异常流；P2=次要功能与边界场景；P3=极端场景、体验类、建议项"
            elif col.maps_to != "custom":
                spec = col.description or canonical_specs.get(col.maps_to, "")
            else:
                spec = (col.description or "按需求内容填写") + f"（自定义字段，值写入 extras[\"{col.name}\"]）"
            suffix = "；必填" if col.required else ""
            if col.maps_to == "custom" and col.enum_values:
                suffix += f"；取值仅允许: {col.enum_values}"
            lines.append(f"- {col.name}: {spec}{suffix}")
        return "\n".join(lines)


class TemplateParseError(ValueError):
    pass


def _normalize(text: str) -> str:
    return re.sub(r"[\s*＊:：()（）]", "", str(text)).lower()


def _map_canonical(header: str) -> CanonicalField:
    normalized = _normalize(header)
    for canonical, patterns in _CANONICAL_PATTERNS:
        if any(p in normalized or normalized in _normalize(p) for p in patterns):
            return canonical
    return "custom"


def _build_template(name: str, headers: list[str], data_rows: list[list[str]]) -> CustomTemplate:
    headers = [str(h).strip() for h in headers if h and str(h).strip()]
    if not headers:
        raise TemplateParseError("模板未识别到表头行（第一行应为字段列名）")

    columns: list[TemplateField] = []
    used: set[str] = set()
    for idx, header in enumerate(headers):
        canonical = _map_canonical(header)
        if canonical in used:  # 同一规范字段只映射一次，重复列视为自定义
            canonical = "custom"
        used.add(canonical)
        samples = [str(r[idx]).strip() for r in data_rows if idx < len(r) and str(r[idx] or "").strip()]
        enum_values = None
        if canonical == "priority":
            enum_values = sorted({s.upper() for s in samples}) if samples else list(PRIORITY_ENUM)
        required = canonical in {"case_id", "module", "title", "priority", "steps", "expected"} or (
            bool(data_rows) and len(samples) == len(data_rows)
        )
        columns.append(
            TemplateField(name=header, maps_to=canonical, required=required, enum_values=enum_values)
        )
    return CustomTemplate(name=name, columns=columns)


def recognize_template(path: str | Path, name: str | None = None) -> CustomTemplate:
    """解析 Excel(.xlsx) / Word(.docx 表格) 模板文件，识别字段结构（F-4-1）。"""
    path = Path(path)
    name = name or path.stem
    if path.suffix.lower() == ".xlsx":
        ws = load_workbook(str(path), read_only=True).active
        rows = [[c if c is not None else "" for c in row] for row in ws.iter_rows(values_only=True)]
        if not rows:
            raise TemplateParseError("模板文件为空")
        return _build_template(name, list(rows[0]), [list(r) for r in rows[1:51]])
    if path.suffix.lower() == ".docx":
        doc = Document(str(path))
        if not doc.tables:
            raise TemplateParseError("Word 模板中未找到表格，请使用包含表头的表格定义模板")
        table = doc.tables[0]
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        data = [[cell.text.strip() for cell in row.cells] for row in table.rows[1:51]]
        return _build_template(name, headers, data)
    raise TemplateParseError(f"不支持的模板格式 {path.suffix}，请上传 .xlsx 或 .docx")
