"""历史用例入库（F-7-2）：Excel/CSV/XMind 存量用例 → 测试用例库。

与导出器（app/exporters）互为镜像：表格列名复用模板识别的启发式映射
（app/templates/custom.py），XMind 按团队模板层级解析（优先级在节点 labels、
前置条件在用例节点 notes）。每条用例渲染为一个独立检索切片。
"""

import csv
import io
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from app.templates.custom import _map_canonical

_PRIORITY_RE = re.compile(r"^[Pp]([0-9])$")
# 禅道 / TestLink 等的数字优先级：1 最高 → P0；文字档位 高/中/低
_NUMERIC_PRIORITY = {"1": "P0", "2": "P1", "3": "P2", "4": "P3"}

# 步骤/预期列中的行前编号："1. xxx" / "1、xxx" / "1) xxx" / 禅道分组子步骤 "1.1 xxx"
_STEP_NO_RE = re.compile(r"^\s*(?:\d+(?:\.\d+)+\s+|\d+\s*[.、)．]\s*)")
_HTML_BR_RE = re.compile(r"<br\s*/?>", re.I)
_HTML_TAG_RE = re.compile(r"<[^>]+>")


def _clean_cell(value) -> str:
    """单元格文本规范化：禅道导出的步骤/预期含 <br /> 换行与 HTML 标签；数字单元格去掉 .0。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value)
    if "<" in text:
        text = _HTML_TAG_RE.sub("", _HTML_BR_RE.sub("\n", text))
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


class CaseImportError(ValueError):
    pass


def normalize_priority(value: str) -> str:
    """优先级归一：大写；P4 及以上并入 P3（2026-07-10 四级决策）；
    数字 1–4（禅道等）对应 P0–P3；其余（含自定义模板的文字档位）原样保留。"""
    value = _clean_cell(value).upper()
    if value in _NUMERIC_PRIORITY:
        return _NUMERIC_PRIORITY[value]
    m = _PRIORITY_RE.match(value)
    if m and int(m.group(1)) > 3:
        return "P3"
    return value


def parse_cases_file(path: str | Path) -> list[dict]:
    """按扩展名解析存量用例文件，返回规范化用例字典列表。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _parse_xlsx(path)
    if suffix == ".csv":
        return _parse_csv(path)
    if suffix == ".xmind":
        return _parse_xmind(path)
    if suffix in (".xls", ".html", ".htm"):
        return _parse_html_table(path)
    raise CaseImportError(f"不支持的用例文件格式 {suffix}，当前支持: .xlsx / .csv / .xmind / 禅道导出的 .xls")


def render_case_chunk(case: dict) -> str:
    """将一条用例渲染为知识切片文本（检索与注入的最小单元）。"""
    head_parts = [f"模块：{case.get('module', '')}"]
    if case.get("priority"):
        head_parts.append(f"优先级：{case['priority']}")
    if case.get("case_id"):
        head_parts.append(f"编号：{case['case_id']}")
    lines = [f"【历史用例】{case.get('title', '')}", " | ".join(head_parts)]
    if case.get("precondition"):
        lines.append(f"前置条件：{case['precondition']}")
    steps = case.get("steps") or []
    if steps:
        lines.append("步骤：")
        lines.extend(
            f"{i}. {s.get('action', '')} → 预期：{s.get('expected', '')}"
            for i, s in enumerate(steps, 1)
        )
    if case.get("keywords"):
        lines.append(f"关键词：{case['keywords']}")
    if case.get("remark"):
        lines.append(f"备注：{case['remark']}")
    return "\n".join(lines)


# ---- 表格（Excel / CSV）----


def _parse_xlsx(path: Path) -> list[dict]:
    from app.parsers.base import check_zip_safety

    check_zip_safety(path)
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = [[_clean_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    return _rows_to_cases(rows, path.name)


def _parse_html_table(path: Path) -> list[dict]:
    """禅道旧版「导出 Excel」得到的 .xls 实为 HTML 表格：解析第一个 <table>，单元格内 <br /> 作换行。"""
    from html.parser import HTMLParser

    from app.parsers.base import read_text_any

    raw = path.read_bytes()
    if raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":  # 真正的二进制 xls（OLE2）
        raise CaseImportError(f"{path.name} 是旧版二进制 Excel（.xls），请在禅道导出时选择 xlsx 或 csv 格式")
    text = read_text_any(path)

    class _Table(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self.rows: list[list[str]] = []
            self._row: list[str] | None = None
            self._cell: list[str] | None = None
            self._done = False

        def handle_starttag(self, tag, attrs):
            if self._done:
                return
            if tag == "tr":
                self._flush_row()
                self._row = []
            elif tag in ("td", "th"):
                if self._row is None:  # 表头单元格不在 <tr> 内的宽松写法
                    self._row = []
                self._cell = []
            elif tag == "br" and self._cell is not None:
                self._cell.append("\n")

        def handle_endtag(self, tag):
            if self._done:
                return
            if tag in ("td", "th") and self._cell is not None and self._row is not None:
                self._row.append("".join(self._cell).strip())
                self._cell = None
            elif tag in ("tr", "thead", "tbody"):
                self._flush_row()
            elif tag == "table":
                self._flush_row()
                if self.rows:
                    self._done = True  # 只取第一张表

        def _flush_row(self):
            if self._row is not None:
                self.rows.append(self._row)
                self._row = None

        def handle_data(self, data):
            if self._cell is not None:
                self._cell.append(data)

    parser = _Table()
    parser.feed(text)
    if not parser.rows:
        raise CaseImportError(f"{path.name} 未找到表格内容（禅道导出请选择 xlsx 或 csv 格式）")
    return _rows_to_cases(parser.rows, path.name)


def _parse_csv(path: Path) -> list[dict]:
    from app.parsers.base import read_text_any

    text = read_text_any(path)
    rows = [[_clean_cell(c) for c in row] for row in csv.reader(io.StringIO(text))]
    return _rows_to_cases(rows, path.name)


def _rows_to_cases(rows: list[list[str]], filename: str) -> list[dict]:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise CaseImportError(f"{filename} 内容为空")
    headers = [str(h).strip() for h in rows[0]]
    mapping = {i: _map_canonical(h) for i, h in enumerate(headers) if h}
    if "title" not in mapping.values():
        raise CaseImportError(
            f"{filename} 表头未识别到用例标题列（首行表头: {[h for h in headers if h]}）"
        )
    cases = []
    for row in rows[1:]:
        record: dict = {}
        for i, canonical in mapping.items():
            if canonical == "custom" or i >= len(row):
                continue
            record.setdefault(canonical, str(row[i] or "").strip())
        if not record.get("title"):
            continue
        record["priority"] = normalize_priority(record.get("priority", ""))
        if record.get("module"):  # 禅道模块导出为「/父模块/子模块」路径：去掉首尾斜杠
            record["module"] = record["module"].strip().strip("/").strip()
        record["steps"] = _pair_steps(record.pop("steps", ""), record.pop("expected", ""))
        cases.append(record)
    if not cases:
        raise CaseImportError(f"{filename} 未解析到任何用例数据行")
    return cases


def _pair_steps(steps_text: str, expected_text: str) -> list[dict]:
    """把「1. 动作」与「1. 预期」两列多行文本还原为步骤对（数量不齐时补空）。"""
    actions = [_STEP_NO_RE.sub("", s).strip() for s in steps_text.splitlines() if s.strip()]
    expecteds = [_STEP_NO_RE.sub("", s).strip() for s in expected_text.splitlines() if s.strip()]
    length = max(len(actions), len(expecteds))
    return [
        {
            "action": actions[i] if i < len(actions) else "",
            "expected": expecteds[i] if i < len(expecteds) else "",
        }
        for i in range(length)
    ]


# ---- XMind（ZEN 格式：zip + content.json）----


def _parse_xmind(path: Path) -> list[dict]:
    from app.parsers.base import check_zip_safety

    check_zip_safety(path)
    try:
        with zipfile.ZipFile(path) as zf:
            sheets = json.loads(zf.read("content.json"))
    except (zipfile.BadZipFile, KeyError, json.JSONDecodeError) as e:
        raise CaseImportError(f"{path.name} 不是有效的 XMind 2020+ 文件（ZEN 格式）: {e}")
    cases: list[dict] = []
    for sheet in sheets:
        root = sheet.get("rootTopic") or {}
        for child in _children(root):
            _walk_topic(child, module_path=[], cases=cases)
    if not cases:
        raise CaseImportError(f"{path.name} 未解析到任何用例节点（用例节点需带优先级标签或备注）")
    return cases


def _children(topic: dict) -> list[dict]:
    return (topic.get("children") or {}).get("attached") or []


def _notes(topic: dict) -> str:
    return (((topic.get("notes") or {}).get("plain") or {}).get("content") or "").strip()


def _priority_label(topic: dict) -> str | None:
    for label in topic.get("labels") or []:
        if _PRIORITY_RE.match(str(label).strip()):
            return normalize_priority(label)
    return None


def _walk_topic(topic: dict, module_path: list[str], cases: list[dict], depth: int = 0) -> None:
    if depth > 64:
        raise CaseImportError("XMind 层级过深（超过 64 级），拒绝导入")
    """递归下钻：带优先级标签或备注的节点为用例；其余带子节点的为模块层级。

    团队模板约定（docs/architecture/测试用例模版.xmind）：优先级在 labels、
    前置条件在用例节点 notes；用例下为「步骤 → 预期结果（步骤子节点）」。
    无标签无备注的叶子节点视为纯测试点（仅标题）。
    """
    title = str(topic.get("title", "")).strip()
    priority = _priority_label(topic)
    notes = _notes(topic)
    children = _children(topic)
    if priority is not None or notes or not children:
        steps = [
            {
                "action": str(step.get("title", "")).strip(),
                "expected": str(_children(step)[0].get("title", "")).strip() if _children(step) else "",
            }
            for step in children
        ]
        cases.append(
            {
                "case_id": "",
                "module": "/".join(module_path) or "未分组",
                "title": title,
                "priority": priority or "",
                "precondition": notes,
                "steps": steps,
                "remark": "",
            }
        )
        return
    for child in children:
        _walk_topic(child, module_path + [title], cases, depth + 1)
